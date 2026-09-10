"""
test_lnk_generation.py — Tests for link model generation pipeline.

Validates:
- _derive_lnk_name naming conventions
- _build_lnk_yaml_model produces correct YAML structure
- LNK XLSX → SQL code generation (end-to-end)
- make_yml produces dbt_constraints.primary_key for links
- validate_tech_spec LNK validation checks
- Integration: parse production lnk_po_item.sql and lnk_bom_plant_item.sql
"""

import sys
import tempfile
from pathlib import Path

import pytest
import yaml

# Add parent dirs so we can import the modules under test
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from pipeline_orchestrator import (
    _derive_lnk_name,
    _build_lnk_yaml_model,
    _parse_existing_hub,
    HUB_WATERMARK_THRESHOLD,
    HUB_FULL_REFRESH_THRESHOLD,
)
from generate_tech_spec import TechSpecBuilder, RAW_VAULT_META_COLUMNS
from make_yml import make_yml


# ---------------------------------------------------------------------------
# _derive_lnk_name tests
# ---------------------------------------------------------------------------

class TestDeriveLnkName:

    def test_simple_name(self):
        assert _derive_lnk_name("po_item") == "lnk_po_item"

    def test_already_prefixed(self):
        assert _derive_lnk_name("lnk_po_item") == "lnk_po_item"

    def test_uppercase_input(self):
        assert _derive_lnk_name("PO_ITEM") == "lnk_po_item"

    def test_whitespace(self):
        assert _derive_lnk_name("  bom_plant_item  ") == "lnk_bom_plant_item"

    def test_already_prefixed_uppercase(self):
        assert _derive_lnk_name("LNK_COPA_SALES") == "lnk_copa_sales"


# ---------------------------------------------------------------------------
# _build_lnk_yaml_model tests
# ---------------------------------------------------------------------------

class TestBuildLnkYamlModel:

    @pytest.fixture
    def state(self):
        return {
            "model_name": "v_psa_stg_po_item__winn_sap",
            "schema": "winn_sap",
            "table": "z_ekpo",
            "bk": "EBELN, EBELP",
            "bk_name": "PO_ITEM_BK",
            "rec_src": "USOHNO.SAP.ECCPRD.Z_EKPO",
        }

    @pytest.fixture
    def profile(self):
        return {
            "ingestion_source": "snp_glue",
            "columns": [
                {"name": "PO_LINE_NUMBER", "type": "TEXT"},
            ],
        }

    @pytest.fixture
    def parent_hks(self):
        return ["PO_ITEM_HK", "PO_HEADER_HK", "ITEM_HK", "SUPPLIER_HK"]

    def test_basic_structure(self, state, profile, parent_hks):
        model = _build_lnk_yaml_model(state, profile, parent_hks, "po_item", [], 1000)
        assert model["layer"] == "LNK"
        assert model["derived_name"] == "lnk_po_item"
        assert model["short_name"] == "PO_ITEM"
        assert len(model["sources"]) == 1

    def test_lhk_naming(self, state, profile, parent_hks):
        model = _build_lnk_yaml_model(state, profile, parent_hks, "po_item", [], 1000)
        lhk_col = model["columns"][0]
        assert lhk_col["staging_column_name"] == "LNK_PO_ITEM_HK"
        assert lhk_col["pk"] == "PK: LNK_PO_ITEM_HK"
        assert lhk_col["ghost_record"] == "hash"

    def test_parent_hks_in_columns(self, state, profile, parent_hks):
        model = _build_lnk_yaml_model(state, profile, parent_hks, "po_item", [], 1000)
        col_names = [c["staging_column_name"] for c in model["columns"]]
        for hk in parent_hks:
            assert hk in col_names
        # All parent HKs should have ghost_record=hash
        for c in model["columns"]:
            if c["staging_column_name"] in parent_hks:
                assert c["ghost_record"] == "hash"

    def test_no_bkcc(self, state, profile, parent_hks):
        model = _build_lnk_yaml_model(state, profile, parent_hks, "po_item", [], 1000)
        col_names = [c["staging_column_name"] for c in model["columns"]]
        assert "BKCC" not in col_names

    def test_no_hashdiff(self, state, profile, parent_hks):
        model = _build_lnk_yaml_model(state, profile, parent_hks, "po_item", [], 1000)
        col_names = [c["staging_column_name"] for c in model["columns"]]
        assert "HASHDIFF" not in col_names

    def test_load_dts_and_rec_src(self, state, profile, parent_hks):
        model = _build_lnk_yaml_model(state, profile, parent_hks, "po_item", [], 1000)
        col_names = [c["staging_column_name"] for c in model["columns"]]
        assert "LOAD_DTS" in col_names
        assert "REC_SRC" in col_names

    def test_dcks_included(self, state, profile, parent_hks):
        model = _build_lnk_yaml_model(
            state, profile, parent_hks, "po_item", ["PO_LINE_NUMBER"], 1000,
        )
        col_names = [c["staging_column_name"] for c in model["columns"]]
        assert "PO_LINE_NUMBER" in col_names
        dck_col = next(c for c in model["columns"] if c["staging_column_name"] == "PO_LINE_NUMBER")
        assert dck_col["ghost_record"] == "value_text"

    def test_dck_datatype_lookup(self, state, profile, parent_hks):
        model = _build_lnk_yaml_model(
            state, profile, parent_hks, "po_item", ["PO_LINE_NUMBER"], 1000,
        )
        dck_col = next(c for c in model["columns"] if c["staging_column_name"] == "PO_LINE_NUMBER")
        assert dck_col["datatype"] == "TEXT"

    def test_column_order(self, state, profile, parent_hks):
        """LNK columns: LHK, parent HKs, DCKs, LOAD_DTS, REC_SRC."""
        model = _build_lnk_yaml_model(
            state, profile, parent_hks, "po_item", ["PO_LINE_NUMBER"], 1000,
        )
        col_names = [c["staging_column_name"] for c in model["columns"]]
        assert col_names[0] == "LNK_PO_ITEM_HK"
        assert col_names[1:5] == ["PO_ITEM_HK", "PO_HEADER_HK", "ITEM_HK", "SUPPLIER_HK"]
        assert col_names[5] == "PO_LINE_NUMBER"
        assert col_names[-2] == "LOAD_DTS"
        assert col_names[-1] == "REC_SRC"

    def test_no_qualify_in_final_filter(self, state, profile, parent_hks):
        """Links do NOT have QUALIFY in FINAL (unlike hubs)."""
        model = _build_lnk_yaml_model(state, profile, parent_hks, "po_item", [], 1000)
        final_filter = model["sources"][0]["final_layer_filter"]
        assert "NOT EXISTS" in final_filter
        assert "qualify" not in final_filter.lower()

    def test_src_qualify_non_watermark(self, state, profile, parent_hks):
        """Non-watermark links have QUALIFY in source_layer_filter."""
        model = _build_lnk_yaml_model(state, profile, parent_hks, "po_item", [], 1000)
        src_filter = model["sources"][0]["source_layer_filter"]
        assert "QUALIFY" in src_filter
        assert "LNK_PO_ITEM_HK" in src_filter

    def test_no_src_qualify_for_watermark(self, state, profile, parent_hks):
        """Watermark links don't have QUALIFY in source_layer_filter (build.py handles it)."""
        model = _build_lnk_yaml_model(
            state, profile, parent_hks, "po_item", [],
            HUB_WATERMARK_THRESHOLD + 1,
        )
        src_filter = model["sources"][0]["source_layer_filter"]
        assert src_filter == ""

    def test_full_refresh_false_config(self, state, profile, parent_hks):
        model = _build_lnk_yaml_model(
            state, profile, parent_hks, "po_item", [],
            HUB_FULL_REFRESH_THRESHOLD + 1,
        )
        model_config = model["sources"][0]["model_config"]
        assert 'full_refresh = var("force_full_refresh", false)' in model_config
        assert "'lnk'" in model_config

    def test_empty_model_config_below_threshold(self, state, profile, parent_hks):
        model = _build_lnk_yaml_model(state, profile, parent_hks, "po_item", [], 1000)
        assert model["sources"][0]["model_config"] == ""

    def test_snp_glue_qualify_order(self, state, profile, parent_hks):
        """LNK always uses LOAD_DTS (reads from staging, not PSA). Lesson #61."""
        model = _build_lnk_yaml_model(state, profile, parent_hks, "po_item", [], 1000)
        assert model["sources"][0]["qualify_order_by"] == "LOAD_DTS"

    def test_fivetran_qualify_order(self, state, parent_hks):
        """LNK always uses LOAD_DTS regardless of ingestion source. Lesson #61."""
        profile = {"ingestion_source": "fivetran", "columns": []}
        model = _build_lnk_yaml_model(state, profile, parent_hks, "po_item", [], 1000)
        assert model["sources"][0]["qualify_order_by"] == "LOAD_DTS"

    def test_source_reads_from_stg(self, state, profile, parent_hks):
        model = _build_lnk_yaml_model(state, profile, parent_hks, "po_item", [], 1000)
        src = model["sources"][0]
        assert src["source_table"] == "v_psa_stg_po_item__winn_sap"
        assert src["source_schema"] == "int_staging_views"
        assert src["target_schema"] == "raw_vault"

    def test_many_parent_hks(self, state, profile):
        """Copa-style link with 13 parent HKs."""
        hks = [
            "COPA_HK", "CONTROLLING_AREA_HK", "COST_CENTER_HK",
            "COST_ELEMENT_HK", "CURRENCY_TYPE_HK", "ORDER_HEADER_HK",
            "ORDER_LINE_HK", "CUSTOMER_HK", "SALES_ORGANIZATION_HK",
            "DISTRIBUTION_CHANNEL_HK", "DIVISION_HK", "ITEM_HK", "PLANT_HK",
        ]
        model = _build_lnk_yaml_model(state, profile, hks, "copa_sales", [], 1000)
        assert model["columns"][0]["staging_column_name"] == "LNK_COPA_SALES_HK"
        # 1 LHK + 13 parent HKs + LOAD_DTS + REC_SRC = 16
        assert len(model["columns"]) == 16

    def test_many_dcks(self, state, profile, parent_hks):
        """Copa-style link with multiple DCKs."""
        dcks = ["MANDT", "VRGAR", "VERSI", "PERIO", "PAOBJNR", "PASUBNR"]
        model = _build_lnk_yaml_model(
            state, profile, parent_hks, "po_item", dcks, 1000,
        )
        col_names = [c["staging_column_name"] for c in model["columns"]]
        for dck in dcks:
            assert dck in col_names
        # All DCKs should have ghost_record=value_text
        for c in model["columns"]:
            if c["staging_column_name"] in dcks:
                assert c["ghost_record"] == "value_text"


# ---------------------------------------------------------------------------
# XLSX Generation (TechSpecBuilder for LNK)
# ---------------------------------------------------------------------------

class TestLnkTechSpec:

    @pytest.fixture
    def lnk_config(self):
        return {
            "schema_version": "1.0",
            "filename": "v_psa_stg_bom_plant_item__winn_sap",
            "_pipeline_metadata": {
                "bkcc_rec_src": "USOHNO.SAP.ECCPRD.STPO",
                "row_count": 1000,
                "volume_tier": "normal",
            },
            "models": [
                {
                    "layer": "LNK",
                    "derived_name": "lnk_bom_plant_item",
                    "short_name": "BOM_PLANT_ITEM",
                    "sources": [{
                        "source_schema": "int_staging_views",
                        "source_table": "v_psa_stg_bom_plant_item__winn_sap",
                        "alias": "SRC",
                        "source_layer_filter": "QUALIFY (ROW_NUMBER() OVER(PARTITION BY LNK_BOM_PLANT_ITEM_HK ORDER BY LOAD_DTS)) = 1",
                        "final_layer_filter": (
                            "{% if is_incremental() %}\n"
                            "WHERE NOT EXISTS (\n"
                            "    SELECT 1\n"
                            "    FROM {{ this }} existing\n"
                            "    WHERE existing.LNK_BOM_PLANT_ITEM_HK = JOIN_RESULT.LNK_BOM_PLANT_ITEM_HK\n"
                            ")\n"
                            "{% endif %}"
                        ),
                        "target_schema": "raw_vault",
                        "model_config": "",
                        "qualify_order_by": "LOAD_DTS",
                    }],
                    "columns": [
                        {"source_table": "SRC", "source_column": "LNK_BOM_PLANT_ITEM_HK",
                         "staging_column_name": "LNK_BOM_PLANT_ITEM_HK", "datatype": "BINARY",
                         "pk": "PK: LNK_BOM_PLANT_ITEM_HK", "ghost_record": "hash"},
                        {"source_table": "SRC", "source_column": "BOM_HK",
                         "staging_column_name": "BOM_HK", "datatype": "BINARY",
                         "ghost_record": "hash"},
                        {"source_table": "SRC", "source_column": "ITEM_HK",
                         "staging_column_name": "ITEM_HK", "datatype": "BINARY",
                         "ghost_record": "hash"},
                        {"source_table": "SRC", "source_column": "PLANT_HK",
                         "staging_column_name": "PLANT_HK", "datatype": "BINARY",
                         "ghost_record": "hash"},
                        {"source_table": "SRC", "source_column": "LOAD_DTS",
                         "staging_column_name": "LOAD_DTS", "datatype": "TIMESTAMP_NTZ",
                         "ghost_record": "load_dts"},
                        {"source_table": "SRC", "source_column": "REC_SRC",
                         "staging_column_name": "REC_SRC", "datatype": "TEXT",
                         "ghost_record": "rec_src"},
                    ],
                },
            ],
        }

    def test_xlsx_creates_tables_and_columns_sheets(self, lnk_config):
        builder = TechSpecBuilder(lnk_config)
        wb = builder.build()
        sheet_names = wb.sheetnames
        assert any("Tables" in s for s in sheet_names)
        assert any("Columns" in s for s in sheet_names)

    def test_xlsx_columns_no_bkcc(self, lnk_config):
        builder = TechSpecBuilder(lnk_config)
        wb = builder.build()
        columns_sheet = next(s for s in wb.sheetnames if "Columns" in s)
        ws = wb[columns_sheet]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        staging_names = [r[8] for r in rows if r[8]]
        assert "BKCC" not in staging_names

    def test_xlsx_columns_no_hashdiff(self, lnk_config):
        builder = TechSpecBuilder(lnk_config)
        wb = builder.build()
        columns_sheet = next(s for s in wb.sheetnames if "Columns" in s)
        ws = wb[columns_sheet]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        staging_names = [r[8] for r in rows if r[8]]
        assert "HASHDIFF" not in staging_names

    def test_xlsx_tables_no_bkcc_row(self, lnk_config):
        """LNK Tables should NOT have a BKCC source row."""
        builder = TechSpecBuilder(lnk_config)
        wb = builder.build()
        tables_sheet = next(s for s in wb.sheetnames if "Tables" in s)
        ws = wb[tables_sheet]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        aliases = [r[2] for r in rows if r[2]]
        assert "ref_bkcc" not in aliases

    def test_lnk_meta_columns_no_psa(self):
        """RAW_VAULT_META_COLUMNS['LNK'] should NOT have PSA_LOAD_DTS or PSA_DELETE_IND."""
        meta_cols = [m["col"] for m in RAW_VAULT_META_COLUMNS["LNK"]]
        assert "PSA_LOAD_DTS" not in meta_cols
        assert "PSA_DELETE_IND" not in meta_cols
        assert "LOAD_DTS" in meta_cols
        assert "REC_SRC" in meta_cols


# ---------------------------------------------------------------------------
# make_yml for LNK
# ---------------------------------------------------------------------------

class TestLnkMakeYml:

    @pytest.fixture
    def lnk_tables(self):
        return {
            "T1": [{
                "DERIVED_NAME": "lnk_po_item",
                "TARGET SCHEMA": "raw_vault",
                "MODEL CONFIG": "",
                "ALIAS": "SRC",
            }],
        }

    @pytest.fixture
    def lnk_columns(self):
        return {
            "C1": {
                "STAGING LAYER COLUMN NAME": "LNK_PO_ITEM_HK",
                "PK": "PK: LNK_PO_ITEM_HK",
                "UNIQUE": "",
                "NOT NULL": "",
                "SOURCE TABLE": "SRC",
                "SOURCE COLUMN": "LNK_PO_ITEM_HK",
                "GHOST RECORD": "hash",
            },
            "C2": {
                "STAGING LAYER COLUMN NAME": "PO_ITEM_HK",
                "PK": "",
                "UNIQUE": "",
                "NOT NULL": "",
                "SOURCE TABLE": "SRC",
                "SOURCE COLUMN": "PO_ITEM_HK",
                "GHOST RECORD": "hash",
            },
            "C3": {
                "STAGING LAYER COLUMN NAME": "SUPPLIER_HK",
                "PK": "",
                "UNIQUE": "",
                "NOT NULL": "",
                "SOURCE TABLE": "SRC",
                "SOURCE COLUMN": "SUPPLIER_HK",
                "GHOST RECORD": "hash",
            },
            "C4": {
                "STAGING LAYER COLUMN NAME": "LOAD_DTS",
                "PK": "",
                "UNIQUE": "",
                "NOT NULL": "",
                "SOURCE TABLE": "SRC",
                "SOURCE COLUMN": "LOAD_DTS",
                "GHOST RECORD": "load_dts",
            },
            "C5": {
                "STAGING LAYER COLUMN NAME": "REC_SRC",
                "PK": "",
                "UNIQUE": "",
                "NOT NULL": "",
                "SOURCE TABLE": "SRC",
                "SOURCE COLUMN": "REC_SRC",
                "GHOST RECORD": "rec_src",
            },
        }

    def test_lnk_severity_warn(self, lnk_tables, lnk_columns):
        result = make_yml("LNK_PO_ITEM", lnk_tables, lnk_columns)
        # Check that tests use severity=warn
        model = result["models"][0]
        for test in model.get("data_tests", []):
            if isinstance(test, dict):
                for key, val in test.items():
                    if isinstance(val, dict) and "config" in val:
                        assert val["config"]["severity"] == "warn"

    def test_lnk_pk_primary_key(self, lnk_tables, lnk_columns):
        result = make_yml("LNK_PO_ITEM", lnk_tables, lnk_columns)
        model = result["models"][0]
        pk_tests = [
            t for t in model.get("data_tests", [])
            if isinstance(t, dict) and "dbt_constraints.primary_key" in t
        ]
        assert len(pk_tests) == 1
        pk_cols = pk_tests[0]["dbt_constraints.primary_key"]["arguments"]["column_names"]
        assert pk_cols == ["LNK_PO_ITEM_HK"]

    def test_lnk_tags_config(self, lnk_tables, lnk_columns):
        result = make_yml("LNK_PO_ITEM", lnk_tables, lnk_columns)
        model = result["models"][0]
        assert model.get("config", {}).get("tags") == ["lnk"]

    def test_lnk_min_row_count(self, lnk_tables, lnk_columns):
        result = make_yml("LNK_PO_ITEM", lnk_tables, lnk_columns)
        model = result["models"][0]
        row_count_tests = [
            t for t in model.get("data_tests", [])
            if isinstance(t, dict)
            and "dbt_expectations.expect_table_row_count_to_be_between" in t
        ]
        assert len(row_count_tests) == 1


# ---------------------------------------------------------------------------
# validate_tech_spec LNK checks
# ---------------------------------------------------------------------------

class TestLnkValidation:

    def _build_lnk_wb(self, columns_data, tables_data=None):
        """Build a minimal XLSX workbook for LNK validation."""
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        # Columns sheet
        ws_cols = wb.create_sheet("LNK BOM_PLANT Columns")
        from generate_tech_spec import HUB_LNK_COLUMNS_HEADERS
        ws_cols.append(HUB_LNK_COLUMNS_HEADERS)
        for row in columns_data:
            ws_cols.append(row)

        # Tables sheet
        ws_tables = wb.create_sheet("LNK BOM_PLANT Tables")
        from generate_tech_spec import TABLES_HEADERS
        ws_tables.append(TABLES_HEADERS)
        if tables_data:
            for row in tables_data:
                ws_tables.append(row)
        else:
            # Default valid row
            ws_tables.append([
                "int_staging_views", "v_psa_stg_bom_plant_item__winn_sap",
                "SRC", "",
                "", "", 1, "", "",
                "{% if is_incremental() %}\nWHERE NOT EXISTS (\n    SELECT 1\n    FROM {{ this }} existing\n    WHERE existing.LNK_BOM_PLANT_ITEM_HK = JOIN_RESULT.LNK_BOM_PLANT_ITEM_HK\n)\n{% endif %}",
                "", "raw_vault", "", "",
            ])

        return wb

    def _validate(self, wb, config=None):
        """Run LNK validation on a workbook."""
        import tempfile
        from validate_tech_spec import TechSpecValidator
        if config is None:
            config = {
                "_pipeline_metadata": {},
                "models": [{
                    "layer": "LNK",
                    "derived_name": "lnk_bom_plant_item",
                    "short_name": "BOM_PLANT",
                    "columns": [],
                }],
            }
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            wb.save(f.name)
            validator = TechSpecValidator(config, f.name)
            return validator.validate()

    def _valid_lnk_columns(self):
        """Return valid LNK column data rows (HUB/LNK format: Ghost Record at idx 13)."""
        return [
            # LHK — PK(10), Unique(11), Not Null(12), Ghost Record(13), Hashdiff(14)
            ["", "SRC", "LNK_BOM_PLANT_ITEM_HK", "BINARY", "", "", "", 1,
             "LNK_BOM_PLANT_ITEM_HK", "BINARY", "PK: LNK_BOM_PLANT_ITEM_HK",
             "", "", "hash", "", "", "", "", "", ""],
            # Parent HK 1
            ["", "SRC", "BOM_HK", "BINARY", "", "", "", 2,
             "BOM_HK", "BINARY", "", "", "", "hash", "", "", "", "", "", ""],
            # Parent HK 2
            ["", "SRC", "ITEM_HK", "BINARY", "", "", "", 3,
             "ITEM_HK", "BINARY", "", "", "", "hash", "", "", "", "", "", ""],
            # Parent HK 3
            ["", "SRC", "PLANT_HK", "BINARY", "", "", "", 4,
             "PLANT_HK", "BINARY", "", "", "", "hash", "", "", "", "", "", ""],
            # LOAD_DTS
            ["", "SRC", "LOAD_DTS", "TIMESTAMP", "", "", "", 5,
             "LOAD_DTS", "TIMESTAMP", "", "", "", "load_dts", "", "", "", "", "", ""],
            # REC_SRC
            ["", "SRC", "REC_SRC", "TEXT", "", "", "", 6,
             "REC_SRC", "TEXT", "", "", "", "rec_src", "", "", "", "", "", ""],
        ]

    def test_valid_lnk_passes(self):
        wb = self._build_lnk_wb(self._valid_lnk_columns())
        errors = self._validate(wb)
        assert len(errors) == 0, f"Unexpected errors: {[str(e) for e in errors]}"

    def test_missing_lhk(self):
        cols = self._valid_lnk_columns()
        # Remove the LHK row (row 0)
        cols = cols[1:]
        wb = self._build_lnk_wb(cols)
        errors = self._validate(wb)
        error_rules = [e.rule for e in errors]
        assert "LNK_LHK_MISSING" in error_rules

    def test_too_few_parent_hks(self):
        # Only 1 parent HK (need >=2)
        cols = [
            ["", "SRC", "LNK_BOM_PLANT_ITEM_HK", "BINARY", "", "", "", 1,
             "LNK_BOM_PLANT_ITEM_HK", "BINARY", "PK: LNK_BOM_PLANT_ITEM_HK",
             "", "", "hash", "", "", "", "", "", ""],
            ["", "SRC", "BOM_HK", "BINARY", "", "", "", 2,
             "BOM_HK", "BINARY", "", "", "", "hash", "", "", "", "", "", ""],
            ["", "SRC", "LOAD_DTS", "TIMESTAMP", "", "", "", 3,
             "LOAD_DTS", "TIMESTAMP", "", "", "", "load_dts", "", "", "", "", "", ""],
            ["", "SRC", "REC_SRC", "TEXT", "", "", "", 4,
             "REC_SRC", "TEXT", "", "", "", "rec_src", "", "", "", "", "", ""],
        ]
        wb = self._build_lnk_wb(cols)
        errors = self._validate(wb)
        error_rules = [e.rule for e in errors]
        assert "LNK_PARENT_HK_COUNT" in error_rules

    def test_has_hashdiff_error(self):
        cols = self._valid_lnk_columns()
        # Add HASHDIFF row (HUB/LNK format: Ghost Record at idx 13)
        cols.append(["", "SRC", "HASHDIFF", "BINARY", "", "", "", 7,
                      "HASHDIFF", "BINARY", "", "", "", "hashdiff", "", "", "", "", "", ""])
        wb = self._build_lnk_wb(cols)
        errors = self._validate(wb)
        error_rules = [e.rule for e in errors]
        assert "LNK_HAS_HASHDIFF" in error_rules

    def test_has_bkcc_error(self):
        cols = self._valid_lnk_columns()
        # Add BKCC row (HUB/LNK format: Ghost Record at idx 13)
        cols.append(["", "SRC", "BKCC", "TEXT", "", "", "", 7,
                      "BKCC", "TEXT", "", "", "", "bkcc", "", "", "", "", "", ""])
        wb = self._build_lnk_wb(cols)
        errors = self._validate(wb)
        error_rules = [e.rule for e in errors]
        assert "LNK_HAS_BKCC" in error_rules

    def test_missing_not_exists_in_tables(self):
        cols = self._valid_lnk_columns()
        tables_data = [[
            "int_staging_views", "v_psa_stg_bom_plant_item__winn_sap",
            "SRC", "", "", "", 1, "", "",
            "",  # Empty final layer filter — missing NOT EXISTS
            "", "raw_vault", "", "",
        ]]
        wb = self._build_lnk_wb(cols, tables_data)
        errors = self._validate(wb)
        error_rules = [e.rule for e in errors]
        assert "LNK_MISSING_NOT_EXISTS" in error_rules


# ---------------------------------------------------------------------------
# Integration: parse production link models
# ---------------------------------------------------------------------------

class TestLnkIntegration:

    @pytest.fixture
    def project_root(self):
        return Path(__file__).resolve().parent.parent.parent.parent

    def test_parse_lnk_bom_plant_item(self, project_root):
        """Parse production lnk_bom_plant_item.sql (2 sources, no DCKs)."""
        lnk_path = project_root / "models" / "raw_vault" / "link" / "lnk_bom_plant_item.sql"
        if not lnk_path.exists():
            pytest.skip(f"Production file not found: {lnk_path}")

        # _parse_existing_hub works for links too (same CTE structure)
        parsed = _parse_existing_hub(str(lnk_path))

        # Should detect 2 sources
        assert len(parsed["sources"]) == 2

        # Should detect LHK
        assert parsed["hk_column"] == "LNK_BOM_PLANT_ITEM_HK"

        # Should detect columns: LHK + 3 parent HKs + LOAD_DTS + REC_SRC
        assert len(parsed["hub_columns"]) >= 5

        # Ghost block should exist
        assert parsed["ghost_record_block"] is not None

    def test_parse_lnk_po_item(self, project_root):
        """Parse production lnk_po_item.sql (7 sources, with DCKs)."""
        lnk_path = project_root / "models" / "raw_vault" / "link" / "lnk_po_item.sql"
        if not lnk_path.exists():
            pytest.skip(f"Production file not found: {lnk_path}")

        parsed = _parse_existing_hub(str(lnk_path))

        # Should detect 7 sources
        assert len(parsed["sources"]) == 7

        # Should detect LHK
        assert parsed["hk_column"] == "LNK_PO_ITEM_HK"

        # Should detect columns including DCK (PO_LINE_NUMBER)
        col_names = [c.upper() for c in parsed["hub_columns"]]
        assert "PO_LINE_NUMBER" in col_names

        # Ghost block should exist
        assert parsed["ghost_record_block"] is not None
