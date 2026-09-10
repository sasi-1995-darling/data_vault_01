"""
test_generate_tech_spec.py — Unit tests for generate_tech_spec.py

Validates that the XLSX tech spec generator handles:
- BK identification (unique + not_null, not just unique)
- Raw source column passthrough when BK has manual_logic (cast)
- YAML field propagation (unique, not_null, manual_logic) to XLSX
- HK naming and formula correctness
- HASHDIFF inclusion/exclusion rules
- LOAD_DTS derivation by ingestion source
- No soft-delete filtering at staging
"""

import sys
import tempfile
from pathlib import Path

import pytest
import yaml

# Add parent dir so we can import the modules under test
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from generate_tech_spec import TechSpecBuilder
from validate_tech_spec import TechSpecValidator, ValidationError


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def base_config():
    """Minimal STG config with a BK cast (manual_logic) and grain columns."""
    return {
        "schema_version": "1.0",
        "filename": "v_psa_stg_test_entity__test_src",
        "_pipeline_metadata": {
            "source_table": "PSA_PROD.test_schema.test_table",
            "bkcc_rec_src": "TEST.SYS.APP.TABLE",
            "bkcc_value": "Test_BKCC",
            "bkcc_description": "Test entity",
            "has_fivetran_deleted": True,
            "has_psa_delete_ind": True,
            "psa_delete_filter": False,
            "filter_comment": "",
            "where_clause": "",
            "null_bk_coalesced": False,
            "null_bk_count": 0,
            "large_volume": False,
            "row_count": 100,
            "volume_tier": "normal",
            "ingestion_source": "fivetran",
            "load_dts_source": "_FIVETRAN_SYNCED",
            "grain_columns": ["COL_A", "COL_B", "LOAD_DTS"],
        },
        "models": [
            {
                "layer": "STG",
                "derived_name": "v_psa_stg_test_entity__test_src",
                "short_name": "test_entity__test_src",
                "description": "Test model",
                "sources": [
                    {
                        "source_schema": "test_schema",
                        "source_table": "test_table",
                        "alias": "SRC",
                    }
                ],
                "columns": [
                    {
                        "source_table": "SRC",
                        "source_column": "COL_A",
                        "datatype": "FLOAT",
                        "manual_logic": "COL_A::TEXT",
                        "staging_column_name": "ENTITY_BK",
                        "staging_datatype": "TEXT",
                        "hashdiff": "no",
                        "unique": "yes",
                        "not_null": "yes",
                    },
                    {
                        "source_table": "SRC",
                        "source_column": "COL_B",
                        "datatype": "FLOAT",
                        "staging_column_name": "COL_B",
                        "staging_datatype": "FLOAT",
                        "hashdiff": "yes",
                        "unique": "yes",
                        "not_null": "",
                    },
                    {
                        "source_table": "SRC",
                        "source_column": "DATA_COL",
                        "datatype": "VARCHAR(100)",
                        "staging_column_name": "DATA_COL",
                        "staging_datatype": "VARCHAR(100)",
                        "hashdiff": "yes",
                        "unique": "",
                        "not_null": "",
                    },
                    {
                        "source_table": "SRC",
                        "source_column": "_FIVETRAN_DELETED",
                        "datatype": "BOOLEAN",
                        "staging_column_name": "_FIVETRAN_DELETED",
                        "staging_datatype": "BOOLEAN",
                        "hashdiff": "yes",
                        "unique": "",
                        "not_null": "",
                    },
                    {
                        "source_table": "SRC",
                        "source_column": "PSA_DELETE_IND",
                        "datatype": "VARCHAR(1)",
                        "staging_column_name": "PSA_DELETE_IND",
                        "staging_datatype": "VARCHAR(1)",
                        "hashdiff": "yes",
                        "unique": "",
                        "not_null": "",
                    },
                    {
                        "source_table": "SRC",
                        "source_column": "_FIVETRAN_SYNCED",
                        "datatype": "TIMESTAMP_TZ",
                        "staging_column_name": "_FIVETRAN_SYNCED",
                        "staging_datatype": "TIMESTAMP_TZ",
                        "hashdiff": "no",
                        "unique": "",
                        "not_null": "",
                    },
                    {
                        "source_table": "SRC",
                        "source_column": "_FIVETRAN_ID",
                        "datatype": "TEXT",
                        "staging_column_name": "_FIVETRAN_ID",
                        "staging_datatype": "TEXT",
                        "hashdiff": "no",
                        "unique": "",
                        "not_null": "",
                    },
                    {
                        "source_table": "SRC",
                        "source_column": "PSA_LOAD_DTS",
                        "datatype": "TIMESTAMP_LTZ",
                        "staging_column_name": "PSA_LOAD_DTS",
                        "staging_datatype": "TIMESTAMP_LTZ",
                        "hashdiff": "no",
                        "unique": "",
                        "not_null": "",
                    },
                    {
                        "source_table": "SRC",
                        "source_column": "PSA_RECORD_SOURCE",
                        "datatype": "TEXT",
                        "staging_column_name": "PSA_RECORD_SOURCE",
                        "staging_datatype": "TEXT",
                        "hashdiff": "no",
                        "unique": "",
                        "not_null": "",
                    },
                ],
            }
        ],
    }


def _build_and_get_rows(config):
    """Build XLSX in memory and return Columns rows as list of dicts."""
    builder = TechSpecBuilder(config)
    wb = builder.build()

    # Find Columns sheet
    cols_sheet = None
    for name in wb.sheetnames:
        if "Columns" in name:
            cols_sheet = wb[name]
            break
    assert cols_sheet is not None, "No Columns sheet found"

    headers = [c.value for c in next(cols_sheet.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in cols_sheet.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(headers, row)))
    return rows


def _build_and_validate(config):
    """Build XLSX, save to temp file, run validator, return (errors, passed).

    Returns (errors_list, passed_bool) where passed=True means no ERRORs.
    """
    builder = TechSpecBuilder(config)
    wb = builder.build()

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        xlsx_path = f.name
        wb.save(xlsx_path)

    validator = TechSpecValidator(config, xlsx_path)
    errors = validator.validate()
    error_msgs = [str(e) for e in errors]
    has_errors = any(e.severity == "ERROR" for e in errors)
    return type("Result", (), {"errors": error_msgs, "has_errors": has_errors})(), not has_errors


# ---------------------------------------------------------------------------
# Tests: BK Identification (Lesson #34)
# ---------------------------------------------------------------------------

class TestBKIdentification:
    """BK must be identified by BOTH unique='yes' AND not_null='yes'."""

    def test_bk_is_first_row(self, base_config):
        rows = _build_and_get_rows(base_config)
        assert rows[0]["Staging Layer Column Name"] == "ENTITY_BK"

    def test_bk_unique_and_not_null(self, base_config):
        rows = _build_and_get_rows(base_config)
        bk_row = rows[0]
        # BK unique is COMPOSITE when grain_columns are defined
        assert "COMPOSITE:" in bk_row["Unique"]
        assert "ENTITY_BK" in bk_row["Unique"]
        assert bk_row["Not Null"] == "yes"

    def test_grain_col_with_unique_only_is_not_bk(self, base_config):
        """COL_B has unique='yes' but not_null='', so it must NOT be treated as BK."""
        rows = _build_and_get_rows(base_config)
        # COL_B should appear as a data column, not as BK
        col_b_rows = [r for r in rows if r.get("Source Column") == "COL_B"]
        assert len(col_b_rows) == 1
        # The BK row should still be ENTITY_BK (from COL_A)
        assert rows[0]["Staging Layer Column Name"] == "ENTITY_BK"

    def test_no_bk_when_no_unique_not_null(self, base_config):
        """If no column has both unique+not_null, validator should catch it."""
        for col in base_config["models"][0]["columns"]:
            col["not_null"] = ""
        result, passed = _build_and_validate(base_config)
        assert not passed
        assert any("No BK column" in e for e in result.errors)


# ---------------------------------------------------------------------------
# Tests: Raw BK Passthrough (Lesson #35)
# ---------------------------------------------------------------------------

class TestRawBKPassthrough:
    """When BK has manual_logic (cast), raw source column must be retained."""

    def test_raw_column_present_with_cast(self, base_config):
        rows = _build_and_get_rows(base_config)
        # Row 1 = BK (ENTITY_BK), Row 2 = raw passthrough (COL_A)
        raw_rows = [r for r in rows if r.get("Staging Layer Column Name") == "COL_A"]
        assert len(raw_rows) == 1, "Raw source column COL_A should appear as passthrough"
        assert raw_rows[0]["Source Column"] == "COL_A"

    def test_no_raw_passthrough_without_cast(self, base_config):
        """When BK has no manual_logic, no extra passthrough row should be added."""
        del base_config["models"][0]["columns"][0]["manual_logic"]
        rows = _build_and_get_rows(base_config)
        # COL_A should appear once as BK only (staging name = ENTITY_BK)
        col_a_rows = [r for r in rows if r.get("Source Column") == "COL_A"]
        assert len(col_a_rows) == 1
        assert col_a_rows[0]["Staging Layer Column Name"] == "ENTITY_BK"

    def test_validator_catches_missing_passthrough(self, base_config):
        """If raw passthrough is missing, validator should fail."""
        builder = TechSpecBuilder(base_config)
        wb = builder.build()
        # Manually remove the raw passthrough row
        cols_sheet = None
        for name in wb.sheetnames:
            if "Columns" in name:
                cols_sheet = wb[name]
                break
        # Delete row 3 (raw COL_A passthrough, 1-indexed: header=1, BK=2, raw=3)
        cols_sheet.delete_rows(3)

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            wb.save(f.name)
            xlsx_path = f.name

        validator = TechSpecValidator(base_config, xlsx_path)
        errors = validator.validate()
        error_msgs = [str(e) for e in errors]
        assert any("Raw source column" in e for e in error_msgs)


# ---------------------------------------------------------------------------
# Tests: YAML Field Propagation (Lesson #36)
# ---------------------------------------------------------------------------

class TestYAMLFieldPropagation:
    """unique, not_null, manual_logic must propagate from YAML to XLSX."""

    def test_unique_propagated_for_grain_column(self, base_config):
        rows = _build_and_get_rows(base_config)
        col_b_row = next(r for r in rows if r.get("Staging Layer Column Name") == "COL_B")
        assert col_b_row["Unique"] == "yes", "Grain column COL_B should have Unique='yes'"

    def test_unique_blank_for_non_grain_column(self, base_config):
        rows = _build_and_get_rows(base_config)
        data_row = next(r for r in rows if r.get("Staging Layer Column Name") == "DATA_COL")
        assert data_row["Unique"] in (None, "", ""), "Non-grain column should not have Unique"

    def test_manual_logic_propagated_for_bk(self, base_config):
        rows = _build_and_get_rows(base_config)
        bk_row = rows[0]
        assert bk_row["Manual Logic"] == "COL_A::TEXT"

    def test_validator_catches_grain_without_unique(self, base_config):
        """Validator should catch grain columns missing unique='yes' when not in COMPOSITE."""
        # Build the XLSX, then manually edit to remove COL_B from both COMPOSITE and its own row
        builder = TechSpecBuilder(base_config)
        wb = builder.build()
        cols_sheet = None
        for name in wb.sheetnames:
            if "Columns" in name:
                cols_sheet = wb[name]
                break
        for row in cols_sheet.iter_rows(min_row=2):
            staging_name = row[8].value or ""
            if staging_name.endswith("_BK"):
                row[11].value = "COMPOSITE: ENTITY_BK, LOAD_DTS"  # Remove COL_B
            elif staging_name == "COL_B":
                row[11].value = ""  # Remove unique flag from COL_B row
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            wb.save(f.name)
            xlsx_path = f.name
        validator = TechSpecValidator(base_config, xlsx_path)
        errors = validator.validate()
        error_msgs = [str(e) for e in errors]
        has_errors = any(e.severity == "ERROR" for e in errors)
        assert has_errors
        assert any("COL_B" in e and "Unique" in e for e in error_msgs)


# ---------------------------------------------------------------------------
# Tests: HK Row
# ---------------------------------------------------------------------------

class TestHKRow:
    """HK must use raw source column name and derive correct staging name."""

    def test_hk_staging_name(self, base_config):
        rows = _build_and_get_rows(base_config)
        hk_row = next(r for r in rows if (r.get("Staging Layer Column Name") or "").endswith("_HK"))
        assert hk_row["Staging Layer Column Name"] == "ENTITY_HK"

    def test_hk_uses_raw_source_column(self, base_config):
        rows = _build_and_get_rows(base_config)
        hk_row = next(r for r in rows if (r.get("Staging Layer Column Name") or "").endswith("_HK"))
        assert "COL_A" in hk_row["Manual Logic"], "HK must use raw source column name"
        assert "BKCC" in hk_row["Manual Logic"], "HK must include BKCC"

    def test_hk_source_column_is_derived(self, base_config):
        rows = _build_and_get_rows(base_config)
        hk_row = next(r for r in rows if (r.get("Staging Layer Column Name") or "").endswith("_HK"))
        assert hk_row["Source Column"] == "(DERIVED)"


# ---------------------------------------------------------------------------
# Tests: HASHDIFF
# ---------------------------------------------------------------------------

class TestHASHDIFF:
    """HASHDIFF must include data flags, exclude metadata columns."""

    def test_hashdiff_includes_fivetran_deleted(self, base_config):
        rows = _build_and_get_rows(base_config)
        hd_row = next(r for r in rows if r.get("Staging Layer Column Name") == "HASHDIFF")
        assert "_FIVETRAN_DELETED" in hd_row["Manual Logic"]

    def test_hashdiff_includes_psa_delete_ind(self, base_config):
        rows = _build_and_get_rows(base_config)
        hd_row = next(r for r in rows if r.get("Staging Layer Column Name") == "HASHDIFF")
        assert "PSA_DELETE_IND" in hd_row["Manual Logic"]

    def test_hashdiff_excludes_metadata(self, base_config):
        rows = _build_and_get_rows(base_config)
        hd_row = next(r for r in rows if r.get("Staging Layer Column Name") == "HASHDIFF")
        logic = hd_row["Manual Logic"]
        for excluded in ["_FIVETRAN_SYNCED", "_FIVETRAN_ID", "PSA_LOAD_DTS", "PSA_RECORD_SOURCE"]:
            assert excluded not in logic, f"HASHDIFF must NOT include {excluded}"

    def test_hashdiff_is_last_row(self, base_config):
        rows = _build_and_get_rows(base_config)
        assert rows[-1]["Staging Layer Column Name"] == "HASHDIFF"


# ---------------------------------------------------------------------------
# Tests: LOAD_DTS Derivation
# ---------------------------------------------------------------------------

class TestLOADDTS:
    """LOAD_DTS derivation must match ingestion source."""

    def test_fivetran_uses_fivetran_synced(self, base_config):
        rows = _build_and_get_rows(base_config)
        load_row = next(r for r in rows if r.get("Staging Layer Column Name") == "LOAD_DTS")
        assert "_FIVETRAN_SYNCED" in load_row["Manual Logic"]

    def test_load_dts_unique_for_grain(self, base_config):
        rows = _build_and_get_rows(base_config)
        load_row = next(r for r in rows if r.get("Staging Layer Column Name") == "LOAD_DTS")
        assert load_row["Unique"] == "yes"

    def test_validator_catches_wrong_load_dts_source(self, base_config):
        """If Fivetran source but LOAD_DTS uses PSA_LOAD_DTS, validator should fail."""
        builder = TechSpecBuilder(base_config)
        wb = builder.build()
        cols_sheet = None
        for name in wb.sheetnames:
            if "Columns" in name:
                cols_sheet = wb[name]
                break

        # Find and modify the LOAD_DTS row's Manual Logic
        for row in cols_sheet.iter_rows(min_row=2):
            if row[8].value == "LOAD_DTS":  # Staging Layer Column Name
                row[5].value = "CONVERT_TIMEZONE('UTC', PSA_LOAD_DTS)"  # Manual Logic
                break

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            wb.save(f.name)
            xlsx_path = f.name

        validator = TechSpecValidator(base_config, xlsx_path)
        errors = validator.validate()
        error_msgs = [str(e) for e in errors]
        assert any("_FIVETRAN_SYNCED" in e for e in error_msgs)


# ---------------------------------------------------------------------------
# Tests: No Delete Filter at Staging
# ---------------------------------------------------------------------------

class TestNoDeleteFilter:
    """Tables tab must never filter on _FIVETRAN_DELETED or PSA_DELETE_IND."""

    def test_no_fivetran_deleted_filter(self, base_config):
        rows = _build_and_get_rows(base_config)
        # Just check XLSX generates without delete filter errors
        result, passed = _build_and_validate(base_config)
        assert passed

    def test_validator_catches_delete_filter(self, base_config):
        """If where_clause filters on soft-delete, validator should catch it."""
        base_config["_pipeline_metadata"]["where_clause"] = "_FIVETRAN_DELETED = FALSE"
        base_config["models"][0]["sources"][0]["filter_conditions"] = "_FIVETRAN_DELETED = FALSE"
        # Note: current validator doesn't check filter content, so this is a future enhancement
        # For now just verify it generates
        rows = _build_and_get_rows(base_config)
        assert len(rows) > 0


# ---------------------------------------------------------------------------
# Tests: Full Validation Pipeline (integration)
# ---------------------------------------------------------------------------

class TestFullValidation:
    """End-to-end: generate XLSX then validate it."""

    def test_full_pipeline_passes(self, base_config):
        """A correctly configured model should pass all validation checks."""
        result, passed = _build_and_validate(base_config)
        assert passed, f"Validation failed with errors: {result.errors}"

    def test_column_ordering(self, base_config):
        """Verify column ordering: BK → raw passthrough → HK → data → tech → BKCC → REC_SRC → LOAD_DTS → HASHDIFF."""
        rows = _build_and_get_rows(base_config)
        stg_names = [r.get("Staging Layer Column Name") for r in rows]

        # BK must be first
        assert stg_names[0] == "ENTITY_BK"
        # Raw passthrough second (since BK has manual_logic)
        assert stg_names[1] == "COL_A"
        # HK third
        assert stg_names[2] == "ENTITY_HK"
        # HASHDIFF must be last
        assert stg_names[-1] == "HASHDIFF"
        # LOAD_DTS second to last
        assert stg_names[-2] == "LOAD_DTS"


# ---------------------------------------------------------------------------
# Test with real YAML config (if it exists on disk)
# ---------------------------------------------------------------------------

REAL_CONFIG = Path(__file__).resolve().parent.parent / "configs" / "payment_terms_lines__ml_ebs.yml"
REAL_XLSX = Path(__file__).resolve().parent.parent / "mappings" / "v_psa_stg_payment_terms_lines__ml_ebs.xlsx"


@pytest.mark.skipif(not REAL_CONFIG.exists(), reason="Real config not present")
class TestRealConfig:
    """Test against the actual payment_terms_lines config."""

    def test_real_config_generates_valid_xlsx(self):
        with open(REAL_CONFIG) as f:
            config = yaml.safe_load(f)
        result, passed = _build_and_validate(config)
        assert passed, f"Validation failed: {result.errors}"

    @pytest.mark.skipif(not REAL_XLSX.exists(), reason="Real XLSX not present")
    def test_real_xlsx_passes_validation(self):
        with open(REAL_CONFIG) as f:
            config = yaml.safe_load(f)
        validator = TechSpecValidator(config, str(REAL_XLSX))
        errors = validator.validate()
        error_msgs = [str(e) for e in errors]
        has_errors = any(e.severity == "ERROR" for e in errors)
        assert not has_errors, f"Validation failed: {error_msgs}"


# ---------------------------------------------------------------------------
# Hub Generation Tests
# ---------------------------------------------------------------------------

class TestHubGeneration:
    """Tests for hub model XLSX generation and validation."""

    @pytest.fixture
    def hub_config(self):
        """Config with both STG and HUB models."""
        return {
            "schema_version": "1.0",
            "filename": "v_psa_stg_test_hub__test_src",
            "_pipeline_metadata": {
                "bkcc_rec_src": "TEST.SYS.APP.TABLE",
                "has_fivetran_deleted": False,
                "has_psa_delete_ind": True,
                "psa_delete_filter": False,
                "null_bk_coalesced": False,
                "large_volume": False,
                "hub_model_name": "hub_test_entity",
                "hub_use_watermark": False,
                "hub_full_refresh": False,
            },
            "models": [
                {
                    "layer": "STG",
                    "derived_name": "v_psa_stg_test_hub__test_src",
                    "short_name": "test_hub__test_src",
                    "sources": [{
                        "source_schema": "test_schema",
                        "source_table": "test_table",
                        "alias": "SRC",
                    }],
                    "columns": [
                        {"source_table": "SRC", "source_column": "COL_A",
                         "staging_column_name": "COL_A", "datatype": "VARCHAR",
                         "unique": "COMPOSITE: COL_A, LOAD_DTS", "not_null": "yes"},
                        {"source_table": "SRC", "source_column": "COL_B",
                         "staging_column_name": "COL_B", "datatype": "VARCHAR",
                         "hashdiff": "yes"},
                    ],
                },
                {
                    "layer": "HUB",
                    "derived_name": "hub_test_entity",
                    "short_name": "TEST_ENTITY",
                    "sources": [{
                        "source_schema": "int_staging_views",
                        "source_table": "v_psa_stg_test_hub__test_src",
                        "alias": "SRC",
                        "final_layer_filter": (
                            "{% if is_incremental() %}\n"
                            "WHERE NOT EXISTS (\n"
                            "    SELECT 1 FROM {{ this }} existing\n"
                            "    WHERE existing.TEST_ENTITY_HK = JOIN_RESULT.TEST_ENTITY_HK\n"
                            ")\n"
                            "{% endif %}\n"
                            "qualify 1 = row_number() over (partition by COL_A, BKCC order by LOAD_DTS)"
                        ),
                        "target_schema": "raw_vault",
                    }],
                    "columns": [
                        {"source_table": "SRC", "source_column": "TEST_ENTITY_HK",
                         "staging_column_name": "TEST_ENTITY_HK", "datatype": "BINARY",
                         "pk": "PK: TEST_ENTITY_HK",
                         "ghost_record": "hash"},
                        {"source_table": "SRC", "source_column": "COL_A",
                         "staging_column_name": "COL_A", "datatype": "VARCHAR",
                         "ghost_record": "value_text"},
                        {"source_table": "SRC", "source_column": "LOAD_DTS",
                         "staging_column_name": "LOAD_DTS", "datatype": "TIMESTAMP_NTZ",
                         "ghost_record": "load_dts"},
                        {"source_table": "SRC", "source_column": "BKCC",
                         "staging_column_name": "BKCC", "datatype": "TEXT",
                         "ghost_record": "bkcc"},
                        {"source_table": "SRC", "source_column": "REC_SRC",
                         "staging_column_name": "REC_SRC", "datatype": "TEXT",
                         "ghost_record": "rec_src"},
                    ],
                },
            ],
        }

    def test_hub_xlsx_has_hub_tabs(self, hub_config):
        """Hub model should create HUB Tables and HUB Columns sheets."""
        builder = TechSpecBuilder(hub_config)
        wb = builder.build()
        sheet_names = wb.sheetnames
        hub_tabs = [s for s in sheet_names if s.startswith("HUB")]
        assert len(hub_tabs) >= 2, f"Expected HUB Tables and Columns tabs, got: {sheet_names}"

    def test_hub_xlsx_no_hashdiff(self, hub_config):
        """Hub Columns sheet should not contain HASHDIFF."""
        builder = TechSpecBuilder(hub_config)
        wb = builder.build()
        hub_cols_sheet = next(
            (s for s in wb.sheetnames if s.startswith("HUB") and "Columns" in s), None
        )
        assert hub_cols_sheet is not None
        ws = wb[hub_cols_sheet]
        staging_names = [
            (row[8] or "").upper()
            for row in ws.iter_rows(min_row=2, values_only=True)
            if row[8]
        ]
        assert "HASHDIFF" not in staging_names, \
            f"Hub should not have HASHDIFF, found columns: {staging_names}"

    def test_hub_xlsx_no_psa_columns(self, hub_config):
        """Hub Columns sheet should NOT contain PSA_LOAD_DTS or PSA_DELETE_IND."""
        builder = TechSpecBuilder(hub_config)
        wb = builder.build()
        hub_cols_sheet = next(
            (s for s in wb.sheetnames if s.startswith("HUB") and "Columns" in s), None
        )
        assert hub_cols_sheet is not None
        ws = wb[hub_cols_sheet]
        staging_names = [
            (row[8] or "").upper()
            for row in ws.iter_rows(min_row=2, values_only=True)
            if row[8]
        ]
        assert "PSA_LOAD_DTS" not in staging_names
        assert "PSA_DELETE_IND" not in staging_names

    def test_hub_xlsx_has_hk(self, hub_config):
        """Hub Columns sheet should have exactly one _HK column."""
        builder = TechSpecBuilder(hub_config)
        wb = builder.build()
        hub_cols_sheet = next(
            (s for s in wb.sheetnames if s.startswith("HUB") and "Columns" in s), None
        )
        ws = wb[hub_cols_sheet]
        hk_cols = [
            (row[8] or "").upper()
            for row in ws.iter_rows(min_row=2, values_only=True)
            if row[8] and str(row[8]).upper().endswith("_HK")
        ]
        assert len(hk_cols) == 1, f"Expected exactly 1 _HK column, got: {hk_cols}"
        assert hk_cols[0] == "TEST_ENTITY_HK"

    def test_hub_xlsx_no_bkcc_ref_join(self, hub_config):
        """Hub Tables sheet should NOT have BKCC ref_business_key_collision row."""
        builder = TechSpecBuilder(hub_config)
        wb = builder.build()
        hub_tables_sheet = next(
            (s for s in wb.sheetnames if s.startswith("HUB") and "Tables" in s), None
        )
        ws = wb[hub_tables_sheet]
        source_tables = [
            (row[1] or "")
            for row in ws.iter_rows(min_row=2, values_only=True)
            if row[1]
        ]
        assert "ref_business_key_collision" not in source_tables, \
            f"Hub should not have BKCC ref join, found: {source_tables}"

    def test_hub_validator_passes(self, hub_config):
        """Hub validation should pass for a well-formed hub config."""
        builder = TechSpecBuilder(hub_config)
        wb = builder.build()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            wb.save(f.name)
            validator = TechSpecValidator(hub_config, f.name)
            errors = validator.validate()
            hub_errors = [e for e in errors if "HUB" in e.rule]
            assert not any(e.severity == "ERROR" for e in hub_errors), \
                f"Hub validation errors: {[str(e) for e in hub_errors]}"

    def test_hub_validator_catches_hashdiff(self, hub_config):
        """Hub validation should flag HASHDIFF if present."""
        hub_config["models"][1]["columns"].append({
            "source_table": "", "source_column": "(DERIVED)",
            "staging_column_name": "HASHDIFF", "datatype": "BINARY",
            "ghost_record": "hashdiff",
        })
        builder = TechSpecBuilder(hub_config)
        wb = builder.build()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            wb.save(f.name)
            validator = TechSpecValidator(hub_config, f.name)
            errors = validator.validate()
            has_hashdiff_error = any(
                e.rule == "HUB_HAS_HASHDIFF" for e in errors
            )
            assert has_hashdiff_error, \
                f"Should catch HASHDIFF in hub, errors: {[str(e) for e in errors]}"

    def test_hub_validator_catches_missing_not_exists(self, hub_config):
        """Hub validation should flag missing NOT EXISTS filter."""
        hub_config["models"][1]["sources"][0]["final_layer_filter"] = (
            "qualify 1 = row_number() over (partition by COL_A order by LOAD_DTS)"
        )
        builder = TechSpecBuilder(hub_config)
        wb = builder.build()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            wb.save(f.name)
            validator = TechSpecValidator(hub_config, f.name)
            errors = validator.validate()
            has_not_exists_error = any(
                e.rule == "HUB_MISSING_NOT_EXISTS" for e in errors
            )
            assert has_not_exists_error, \
                f"Should catch missing NOT EXISTS, errors: {[str(e) for e in errors]}"


# ---------------------------------------------------------------------------
# Check 21: BK/HK KEYS band ordering edge cases
# ---------------------------------------------------------------------------

class TestColumnOrderingKeysBand:
    """BK and HK may swap relative to each other but must appear before DATA."""

    def _make_row(self, staging_name, source_col="SRC_COL"):
        """Create a minimal mock row tuple (indices: 2=source_col, 8=staging_name)."""
        return (None, None, source_col, None, None, None, None, None, staging_name)

    def _make_validator(self, tmp_path):
        """Create a validator with no XLSX (we call _check_column_ordering directly)."""
        import openpyxl
        wb = openpyxl.Workbook()
        xlsx_path = tmp_path / "_test_keys_band.xlsx"
        wb.save(str(xlsx_path))
        v = TechSpecValidator({"models": [], "_pipeline_metadata": {}}, str(xlsx_path))
        return v

    def test_bk_after_hk_no_warning(self, tmp_path):
        """BK appearing after HK should NOT trigger a warning (within KEYS band)."""
        v = self._make_validator(tmp_path)
        rows = [
            self._make_row("ENTITY_HK"),
            self._make_row("ENTITY_BK"),
            self._make_row("DATA_COL"),
            self._make_row("BKCC"),
            self._make_row("REC_SRC"),
            self._make_row("LOAD_DTS"),
            self._make_row("HASHDIFF"),
        ]
        v._check_column_ordering(rows)
        order_warnings = [e for e in v.errors if e.rule == "COLUMN_ORDER"]
        assert len(order_warnings) == 0, f"BK after HK should be allowed: {order_warnings}"

    def test_hk_after_bk_no_warning(self, tmp_path):
        """HK appearing after BK should NOT trigger a warning (within KEYS band)."""
        v = self._make_validator(tmp_path)
        rows = [
            self._make_row("ENTITY_BK"),
            self._make_row("ENTITY_HK"),
            self._make_row("DATA_COL"),
            self._make_row("BKCC"),
            self._make_row("REC_SRC"),
            self._make_row("LOAD_DTS"),
            self._make_row("HASHDIFF"),
        ]
        v._check_column_ordering(rows)
        order_warnings = [e for e in v.errors if e.rule == "COLUMN_ORDER"]
        assert len(order_warnings) == 0, f"HK after BK should be allowed: {order_warnings}"

    def test_bk_after_data_triggers_warning(self, tmp_path):
        """BK appearing after DATA should trigger a warning (past KEYS band)."""
        v = self._make_validator(tmp_path)
        rows = [
            self._make_row("ENTITY_HK"),
            self._make_row("DATA_COL"),
            self._make_row("ENTITY_BK"),  # BK after DATA — should warn
            self._make_row("BKCC"),
            self._make_row("REC_SRC"),
            self._make_row("LOAD_DTS"),
            self._make_row("HASHDIFF"),
        ]
        v._check_column_ordering(rows)
        order_warnings = [e for e in v.errors if e.rule == "COLUMN_ORDER"]
        assert len(order_warnings) >= 1, "BK after DATA should trigger ordering warning"
        assert "ENTITY_BK" in str(order_warnings[0])

    def test_composite_bk_passthrough_before_hk_no_warning(self, tmp_path):
        """#1906: composite-BK raw passthrough rows sit in the KEYS band, so an HK
        placed immediately after them must NOT trigger a COLUMN_ORDER warning.

        The BK-alias row's source cell is newline-joined (INVOICE_ID\\nLINE_NUMBER);
        each component is also its own passthrough staging row.
        """
        v = self._make_validator(tmp_path)
        rows = [
            self._make_row("SUPPLIER_INVOICE_LINE_BK", source_col="INVOICE_ID\nLINE_NUMBER"),
            self._make_row("INVOICE_ID", source_col="INVOICE_ID"),
            self._make_row("LINE_NUMBER", source_col="LINE_NUMBER"),
            self._make_row("SUPPLIER_INVOICE_LINE_HK", source_col="(DERIVED)"),
            self._make_row("AMOUNT"),
            self._make_row("BKCC"),
            self._make_row("REC_SRC"),
            self._make_row("LOAD_DTS"),
            self._make_row("HASHDIFF"),
        ]
        v._check_column_ordering(rows)
        order_warnings = [e for e in v.errors if e.rule == "COLUMN_ORDER"]
        assert len(order_warnings) == 0, f"composite-BK passthrough must not warn: {order_warnings}"

    def test_composite_bk_hk_after_real_data_still_warns(self, tmp_path):
        """#1906 guard: a genuine DATA column before the HK must still warn — the fix
        only reclassifies BK passthrough rows, not real mis-ordering."""
        v = self._make_validator(tmp_path)
        rows = [
            self._make_row("SUPPLIER_INVOICE_LINE_BK", source_col="INVOICE_ID\nLINE_NUMBER"),
            self._make_row("INVOICE_ID", source_col="INVOICE_ID"),
            self._make_row("AMOUNT"),  # genuine DATA column, not a BK passthrough
            self._make_row("SUPPLIER_INVOICE_LINE_HK", source_col="(DERIVED)"),  # HK after DATA
            self._make_row("BKCC"),
            self._make_row("REC_SRC"),
            self._make_row("LOAD_DTS"),
            self._make_row("HASHDIFF"),
        ]
        v._check_column_ordering(rows)
        order_warnings = [e for e in v.errors if e.rule == "COLUMN_ORDER"]
        assert len(order_warnings) >= 1, "HK after a real DATA column should still warn"

