"""
generate_tech_spec.py — Create XLSX tech-spec workbooks for the process-mapping tool.

This script generates properly formatted XLSX files that the process-mapping
tool (src/main.py) can consume to produce uniform dbt SQL + YAML.

Usage:
    python scripts/generate_tech_spec.py --config <config.yml> --outdir <path>/mappings/

Config YAML format:
    See scripts/configs/ for examples (e.g., product_cost_estimate.yml).
"""

import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment
import yaml

# Shared source_column helpers (single source of truth — src/source_column_utils.py)
sys.path.insert(0, str(Path(__file__).resolve().parent / "src"))
from source_column_utils import get_source_columns, source_columns_to_str  # noqa: E402

# Cell style used when a Source Column cell holds a multi-column (composite) key —
# the tokens are newline-separated and wrapped so a composite is never rendered as
# an indistinguishable comma-joined single cell.
_WRAP_ALIGNMENT = Alignment(wrap_text=True, vertical="top")


def _source_column_cell(col):
    """Render a column's source_column for the XLSX 'Source Column' cell.

    Single-column -> the bare token. Composite -> newline-separated tokens so the
    multiplicity is visually and machine-unambiguous (never comma-joined). Order
    is preserved.
    """
    return "\n".join(get_source_columns(col))


def _append_row(ws, row):
    """Append a row and wrap any cell that contains newline-separated tokens."""
    ws.append(row)
    r = ws.max_row
    for c_idx, val in enumerate(row, start=1):
        if isinstance(val, str) and "\n" in val:
            ws.cell(row=r, column=c_idx).alignment = _WRAP_ALIGNMENT



# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

TABLES_HEADERS = [
    "Source Schema", "Source Table", "Alias", "Source Layer Filter",
    "Filter Conditions", "Filter Restriction Rule", "Parent Join Number",
    "Parent Table Join", "Child Table Join", "Final Layer Filter",
    "Join Type", "Target Schema", "Model Config", "Qualify Order By"
]

COLUMNS_HEADERS = [
    "Source Schema", "Source Table", "Source Column", "Datatype",
    "Automated Logic", "Manual Logic", "Mapping Notes", "Order#",
    "Staging Layer Column Name", "Staging Layer Datatype", "PK",
    "Unique", "Not Null", "Hashdiff", "Ghost Record", "Remove Column",
    "Set Default", "Test Expression", "Accepted Values", "Relationship"
]

# HUB/LNK production format: PK → Unique → Not Null → Ghost Record → Hashdiff (20 cols)
HUB_LNK_COLUMNS_HEADERS = [
    "Source Schema", "Source Table", "Source Column", "Datatype",
    "Automated Logic", "Manual Logic", "Mapping Notes", "Order#",
    "Staging Layer Column Name", "Staging Layer Datatype", "PK",
    "Unique", "Not Null", "Ghost Record", "Hashdiff", "Remove Column",
    "Set Default", "Test Expression", "Accepted Values", "Relationship"
]

# SAT production format: PK → Unique → Ghost Record → Not Null (19 cols, no Hashdiff header)
SAT_COLUMNS_HEADERS = [
    "Source Schema", "Source Table", "Source Column", "Datatype",
    "Automated Logic", "Manual Logic", "Mapping Notes", "Order#",
    "Staging Layer Column Name", "Staging Layer Datatype", "PK",
    "Unique", "Ghost Record", "Not Null", "Remove Column",
    "Set Default", "Test Expression", "Accepted Values", "Relationship"
]

# Standard DV meta columns added to every RAW_VAULT model
RAW_VAULT_META_COLUMNS = {
    "HUB": [
        {"col": "BKCC", "dtype": "TEXT", "ghost": "bkcc"},
        {"col": "LOAD_DTS", "dtype": "TIMESTAMP", "ghost": "load_dts"},
        {"col": "REC_SRC", "dtype": "TEXT", "ghost": "rec_src"},
    ],
    "SAT": [
        {"col": "HASHDIFF", "dtype": "BINARY", "ghost": "hashdiff"},
        {"col": "LOAD_DTS", "dtype": "TIMESTAMP", "ghost": "load_dts"},
        {"col": "REC_SRC", "dtype": "TEXT", "ghost": "rec_src"},
        {"col": "PSA_LOAD_DTS", "dtype": "TIMESTAMP", "ghost": "psa_load_dts"},
        {"col": "PSA_DELETE_IND", "dtype": "TEXT", "ghost": "psa_delete_ind"},
    ],
    "LNK": [
        {"col": "LOAD_DTS", "dtype": "TIMESTAMP", "ghost": "load_dts"},
        {"col": "REC_SRC", "dtype": "TEXT", "ghost": "rec_src"},
    ],
}
RAW_VAULT_META_COLUMNS["LSAT"] = RAW_VAULT_META_COLUMNS["SAT"]
RAW_VAULT_META_COLUMNS["MSAT"] = RAW_VAULT_META_COLUMNS["SAT"]

INCREMENTAL_SRC_FILTER = (
    "{% if is_incremental() %}\n"
    " WHERE SRC.LOAD_DTS > (SELECT DATEADD('HOUR', -1, MAX(LOAD_DTS)) FROM {{this}})\n"
    " {% endif %}"
)


def build_incremental_final_filter(hk_col, extra_cols=None):
    """Build the standard incremental NOT EXISTS block for the Final Layer Filter."""
    lines = [
        "{% if is_incremental() %}",
        "WHERE NOT EXISTS (",
        "    SELECT 1 ",
        "    FROM {{ this }} existing",
        f"    WHERE existing.{hk_col} = JOIN_RESULT.{hk_col}",
    ]
    for c in (extra_cols or []):
        lines.append(f"    AND   existing.{c} = JOIN_RESULT.{c}")
    lines += [")", "{% endif %}"]
    return "\n".join(lines)


def build_qualify_block(pk_cols):
    """Build the qualify dedup block appended to Final Layer Filter for HUBs."""
    cols = ", ".join(pk_cols)
    return (
        f"/* The following qualifier is implemented to prevent multiple loads of "
        f"touched records during the initial build, such as multiple rows per BKs "
        f"and BKCC. */\n"
        f"qualify 1 = row_number() over (partition by {cols} order by LOAD_DTS)"
    )


# ---------------------------------------------------------------------------
# Workbook Builder
# ---------------------------------------------------------------------------

class TechSpecBuilder:
    """Builds an XLSX tech-spec workbook from a config dict."""

    def __init__(self, config: dict):
        self.config = config
        self.metadata = config.get("_pipeline_metadata", {})
        self.wb = openpyxl.Workbook()
        # Remove default sheet
        self.wb.remove(self.wb.active)
        self.index_entries = []

    def build(self):
        used_sheet_names = set()
        for model in self.config.get("models", []):
            layer = model["layer"]
            short_name = model["short_name"]
            derived_name = model.get("derived_name", f"{layer.lower()}_{short_name.replace(' ', '_').lower()}")

            tables_tab = self._safe_sheet_name(f"{layer} {short_name} Tables")
            columns_tab = self._safe_sheet_name(f"{layer} {short_name} Columns")

            # Uniqueness guard: append counter suffix if name collides after truncation
            tables_tab = self._unique_sheet_name(tables_tab, used_sheet_names)
            columns_tab = self._unique_sheet_name(columns_tab, used_sheet_names)

            self._add_tables_sheet(tables_tab, model)
            self._add_columns_sheet(columns_tab, model, layer)

            self.index_entries.append((derived_name, tables_tab, columns_tab))

        self._add_index_sheet()
        return self.wb

    @staticmethod
    def _unique_sheet_name(name: str, used: set, max_len: int = 31) -> str:
        """Ensure sheet name is unique by inserting a counter before the suffix.

        Preserves the ' Tables' / ' Columns' suffix so the external code
        generator can match sheet pairs.  Uses max_len - 8 (for ' Columns')
        minus counter length as the base budget, keeping Tables/Columns
        bases identical.
        """
        if name not in used:
            used.add(name)
            return name
        # Split into base + type-suffix (' Tables' or ' Columns')
        suffix_start = name.rfind(" ")
        if suffix_start == -1:
            # No space — fall back to simple append
            for i in range(2, 100):
                counter = f"_{i}"
                candidate = name[:max_len - len(counter)] + counter
                if candidate not in used:
                    used.add(candidate)
                    return candidate
            raise ValueError(f"Cannot create unique sheet name for \'{name}\'")
        type_suffix = name[suffix_start:]   # ' Tables' or ' Columns'
        base = name[:suffix_start]
        for i in range(2, 100):
            counter = f"_{i}"
            # Reserve 8 chars for longest suffix + counter length
            max_base = max_len - 8 - len(counter)
            candidate = base[:max_base] + counter + type_suffix
            if len(candidate) <= max_len and candidate not in used:
                used.add(candidate)
                return candidate
        raise ValueError(f"Cannot create unique sheet name for \'{name}\'")

    @staticmethod
    def _safe_sheet_name(name: str, max_len: int = 31) -> str:
        """Truncate sheet name to Excel's 31-character limit.

        ALWAYS reserves 8 chars for the longest suffix (' Columns')
        so that Tables and Columns tabs share the same truncated base,
        even when the Tables variant (7-char suffix) would fit at 31.
        """
        suffix_start = name.rfind(" ")
        if suffix_start == -1:
            return name[:max_len]
        suffix = name[suffix_start:]  # e.g., ' Tables' or ' Columns'
        base = name[:suffix_start]
        # Always reserve 8 chars (len(' Columns')) regardless of current suffix
        max_base_len = max_len - 8
        return base[:max_base_len] + suffix

    # --- Index ---
    def _add_index_sheet(self):
        ws = self.wb.create_sheet("Index", 0)
        ws.append(["INDEX"])
        for derived, tables_tab, columns_tab in self.index_entries:
            row_num = ws.max_row + 1
            cell_tables = ws.cell(row=row_num, column=1, value=f"{derived} Tables")
            cell_tables.hyperlink = f"#'{tables_tab}'!A1"
            cell_tables.style = "Hyperlink"

            row_num = ws.max_row + 1
            cell_cols = ws.cell(row=row_num, column=1, value=f"{derived} Columns")
            cell_cols.hyperlink = f"#'{columns_tab}'!A1"
            cell_cols.style = "Hyperlink"

    # --- Tables ---
    def _add_tables_sheet(self, tab_name, model):
        ws = self.wb.create_sheet(tab_name)
        ws.append(TABLES_HEADERS)

        for src in model.get("sources", []):
            source_layer_filter = src.get("source_layer_filter", "")
            # Auto-inject QUALIFY dedup for Fivetran driver tables (STG layer only)
            # But ONLY when grain is NOT valid — valid grain means no dedup needed (lesson #82)
            layer = model.get("layer", "").upper()
            alias = src.get("alias", "")
            grain_valid = self.metadata.get("grain_valid", True)  # default True: assume valid unless profile says otherwise
            if layer == "STG" and alias == "SRC" and not source_layer_filter and not grain_valid:
                load_dts_col = self.metadata.get("load_dts_derivation", "")
                if "_FIVETRAN_SYNCED" in load_dts_col.upper():
                    # Find the first BK raw source column for PARTITION BY
                    bk_raw_cols = []
                    for col in model.get("columns", []):
                        ml = col.get("manual_logic", "")
                        sc_list = get_source_columns(col)
                        sn = (col.get("staging_column_name", "") or "").upper()
                        if sn.endswith("_BK"):
                            import re as _re
                            if sc_list == ["(DERIVED)"] and ml:
                                # Extract raw col from TO_CHAR(RAW_COL) or CAST patterns
                                m = _re.search(r"TO_CHAR\((\w+)\)", ml)
                                if m:
                                    bk_raw_cols.append(m.group(1))
                                else:
                                    # Try ::TEXT cast: "COL_NAME::TEXT" -> COL_NAME
                                    m2 = _re.match(r"^(\w+)::", ml)
                                    if m2:
                                        bk_raw_cols.append(m2.group(1))
                            elif sc_list and sc_list != ["(DERIVED)"]:
                                # Direct column BK: source_column holds the raw col(s)
                                bk_raw_cols.extend(sc_list)
                    if bk_raw_cols:
                        # PARTITION BY: ALL BK raw cols + LOAD_DTS source col (_FIVETRAN_SYNCED)
                        # ORDER BY: PSA_DELETE_IND first (truncate-and-load), then _FIVETRAN_SYNCED DESC
                        # Per qualify-dedup-patterns.md and Lesson #65
                        partition_cols = bk_raw_cols + ["_FIVETRAN_SYNCED"]
                        partition_str = ", ".join(partition_cols)
                        source_layer_filter = f"QUALIFY ROW_NUMBER() OVER(PARTITION BY {partition_str} ORDER BY PSA_DELETE_IND, _FIVETRAN_SYNCED DESC) = 1"
            row = [
                src.get("source_schema", ""),
                src.get("source_table", ""),
                src.get("alias", ""),
                source_layer_filter,
                src.get("filter_conditions", self.metadata.get("where_clause", "")),  # Apply where_clause if defined
                src.get("filter_restriction_rule", ""),
                src.get("parent_join_number", 1),
                src.get("parent_table_join", ""),
                src.get("child_table_join", ""),
                src.get("final_layer_filter", ""),
                src.get("join_type", ""),
                src.get("target_schema", ""),
                src.get("model_config", ""),
                src.get("qualify_order_by", ""),
            ]
            ws.append(row)

        # Auto-inject BKCC source row when bkcc_rec_src is configured (STG only).
        # Hub/Link models read BKCC directly from their upstream v_psa_stg source.
        layer = model.get("layer", "").upper()
        bkcc_rec_src = self.metadata.get("bkcc_rec_src", "")
        if bkcc_rec_src and layer == "STG":
            bkcc_row = [
                "raw_vault",                              # Source Schema
                "ref_business_key_collision",              # Source Table
                "ref_bkcc",                               # Alias (descriptive, not just 'A')
                f"WHERE rec_src = '{bkcc_rec_src}'",       # Source Layer Filter (WHERE required — build.py renders verbatim)
                "",                                       # Filter Conditions
                "",                                       # Filter Restriction Rule
                2,                                        # Parent Join Number (secondary join)
                "'1'",                                    # Parent Table Join
                "'1'",                                    # Child Table Join
                "",                                       # Final Layer Filter
                "inner join",                             # Join Type
                "",                                       # Target Schema
                "",                                       # Model Config
                "",                                       # Qualify Order By
            ]
            ws.append(bkcc_row)

    # --- Columns ---
    def _add_columns_sheet(self, tab_name, model, layer):
        ws = self.wb.create_sheet(tab_name)
        # Select correct column headers based on layer
        layer_upper = layer.upper()
        if layer_upper in ("HUB", "LNK"):
            ws.append(HUB_LNK_COLUMNS_HEADERS)
        elif layer_upper in ("SAT", "LSAT", "MSAT", "LMSAT", "ESAT"):
            ws.append(SAT_COLUMNS_HEADERS)
        else:
            ws.append(COLUMNS_HEADERS)

        if layer == "STG":
            # For STG models, orchestrate column order: BK, HK, data, technical, BKCC, REC_SRC, LOAD_DTS, HASHDIFF
            self._add_stg_columns_sheet(ws, model)
        else:
            # For RAW_VAULT models, use original logic
            order = 1
            for col in model.get("columns", []):
                row = self._make_column_row(col, order, layer=layer)
                ws.append(row)
                order += 1

            # Add standard meta columns for RAW_VAULT layers
            meta_key = layer if layer in RAW_VAULT_META_COLUMNS else None
            if layer.startswith("LSAT"):
                meta_key = "LSAT"
            elif layer.startswith("MSAT"):
                meta_key = "MSAT"
            elif layer.startswith("SAT"):
                meta_key = "SAT"

            if meta_key and meta_key in RAW_VAULT_META_COLUMNS:
                first_alias = model["sources"][0].get("alias", "") if model.get("sources") else ""
                for meta in RAW_VAULT_META_COLUMNS[meta_key]:
                    # Skip if user already defined this column (or a named variant like HASHDIFF_*)
                    existing = [c.get("staging_column_name", "").upper() for c in model.get("columns", [])]
                    if meta["col"].upper() in existing:
                        continue
                    # Skip HASHDIFF if a named variant (HASHDIFF_*) already exists (per-SAT HASHDIFF)
                    if meta["col"].upper() == "HASHDIFF" and any(e.startswith("HASHDIFF_") for e in existing):
                        continue
                    meta_row = self._make_meta_row(
                        first_alias, meta["col"], meta["dtype"],
                        meta["ghost"], order, layer=layer,
                    )
                    ws.append(meta_row)
                    order += 1

    def _add_stg_columns_sheet(self, ws, model):
        """Generate Columns sheet for STG model with proper ordering."""
        source_schema = model["sources"][0].get("source_schema", "") if model.get("sources") else ""
        driver_alias = model["sources"][0].get("alias", "") if model.get("sources") else ""
        bkcc_rec_src = self.metadata.get("bkcc_rec_src", "")
        bkcc_value = self.metadata.get("bkcc_value", "[BKCC_VALUE]")
        order = 1

        # Build set of existing staging column names (to prevent duplicates)
        existing_staging_cols = {col.get("staging_column_name", "").upper() for col in model.get("columns", [])}

        # Find BK and HK columns
        bk_col = None
        hk_cols = []  # Collect ALL HK columns
        raw_bk_col = None
        data_cols = []
        technical_cols = []
        yaml_reserved = {}  # YAML-defined reserved columns (BKCC, REC_SRC, LOAD_DTS, HASHDIFF)

        reserved_staging_names = {"BKCC", "REC_SRC", "LOAD_DTS", "HASHDIFF"}  # These go at the end

        for col in model.get("columns", []):
            staging_name = col.get("staging_column_name", "").upper()
            unique_val = str(col.get("unique", "")).strip()
            not_null_val = str(col.get("not_null", "")).strip().lower()
            # BK: has unique (any truthy value) AND not_null=yes
            if unique_val and not_null_val == "yes" and staging_name.endswith("_BK"):
                bk_col = col
            elif staging_name.endswith("_HK") and get_source_columns(col) == ["(DERIVED)"]:
                hk_cols.append(col)
            elif staging_name in reserved_staging_names:
                # Stash YAML-defined reserved columns — they'll be added at the end
                yaml_reserved[staging_name] = col
                continue
            elif col.get("staging_column_name", "").startswith("_FIVETRAN") or col.get("staging_column_name", "").startswith("PSA_"):
                technical_cols.append(col)
            else:
                data_cols.append(col)

        # Extract raw BK passthrough from data_cols — place after BK, before HK.
        # When the orchestrator pre-adds the raw source column to YAML alongside
        # the BK alias, it lands in data_cols. Extract it so it outputs in the
        # BK section (correct order: BK → raw BK → HK → data).
        if bk_col:
            bk_source_cols = get_source_columns(bk_col)
            bk_staging_upper = bk_col.get("staging_column_name", "").upper()
            for i, c in enumerate(data_cols):
                # The raw passthrough shares the BK's source column(s) but keeps the
                # raw column name(s) as its staging name (not the BK alias).
                if (get_source_columns(c) == bk_source_cols
                        and bk_source_cols != ["(DERIVED)"]
                        and c.get("staging_column_name", "").upper() != bk_staging_upper):
                    raw_bk_col = data_cols.pop(i)
                    break

        # Also filter out HK columns from data_cols (handled separately)
        data_cols = [c for c in data_cols if not c.get("staging_column_name", "").upper().endswith("_HK")]

        # Build the COMPOSITE grain string from metadata or BK column
        # Filter out: (1) raw BK source cols (already represented by BK alias),
        # (2) ingestion-specific LOAD_DTS source cols (LOAD_DTS always appended as constant)
        _LOAD_DTS_SOURCE_COLS = {"PSA_LOAD_DTS", "_FIVETRAN_SYNCED", "GLCHANGETIME", "SNP_LOAD_DTS", "LOAD_DTS"}
        grain_columns = self.metadata.get("grain_columns", [])
        if grain_columns:
            bk_staging = bk_col.get("staging_column_name", "") if bk_col else ""
            bk_source_cols = {c.upper() for c in get_source_columns(bk_col)} if bk_col else set()
            composite_parts = [bk_staging]  # Always start with BK alias
            for gc in grain_columns:
                gc_upper = gc.upper()
                # Skip raw BK source col (already represented by BK alias)
                if gc_upper in bk_source_cols:
                    continue
                # Skip ingestion-specific LOAD_DTS columns
                if gc_upper in _LOAD_DTS_SOURCE_COLS:
                    continue
                composite_parts.append(gc_upper)
            # Always end with LOAD_DTS
            composite_parts.append("LOAD_DTS")
            bk_unique_str = f"COMPOSITE: {', '.join(composite_parts)}"
        elif bk_col:
            bk_staging = bk_col.get("staging_column_name", "")
            bk_unique_str = f"COMPOSITE: {bk_staging}, LOAD_DTS"
        else:
            bk_unique_str = "yes"

        # 1. Add BK column first
        if bk_col:
            bk_unique = str(bk_col.get("unique", bk_unique_str))
            # If YAML just says "yes", upgrade to COMPOSITE form
            if bk_unique.lower() == "yes":
                bk_unique = bk_unique_str
            row = [
                source_schema,                              # Source Schema
                driver_alias,                                # Source Table (alias)
                _source_column_cell(bk_col),                # Source Column (composite -> newline-split)
                bk_col.get("datatype", ""),                 # Datatype
                "",                                         # Automated Logic
                bk_col.get("manual_logic", ""),             # Manual Logic (e.g. TERM_ID::TEXT for cast BKs)
                "",                                         # Mapping Notes
                order,                                      # Order#
                bk_col.get("staging_column_name", ""),      # Staging Layer Column Name
                bk_col.get("staging_datatype", ""),         # Staging Layer Datatype
                "",                                         # PK
                bk_unique,                                  # Unique (COMPOSITE: BK, LOAD_DTS)
                "yes",                                      # Not Null
                "",                                         # Hashdiff
                "",                                         # Ghost Record
                "",                                         # Remove Column
                "",                                         # Set Default
                "",                                         # Test Expression
                "",                                         # Accepted Values
                "",                                         # Relationship
            ]
            _append_row(ws, row)
            order += 1

            # If BK has manual_logic (cast), add raw source column as passthrough
            # BUT only if the orchestrator didn't already include it in the YAML columns.
            # Single-column cast BKs only (TERM_ID::TEXT -> TERM_ID). Composite BKs
            # already carry a raw passthrough entry (handled via raw_bk_col below).
            bk_src_cols = get_source_columns(bk_col)
            if (bk_col.get("manual_logic") and bk_src_cols != ["(DERIVED)"]
                    and len(bk_src_cols) == 1):
                # Raw passthrough column name (already cast-stripped in source_column)
                raw_source = bk_src_cols[0]
                # Check if raw passthrough already exists in YAML columns
                already_has_raw = any(
                    get_source_columns(c) == [raw_source]
                    and c.get("staging_column_name", "").upper() == raw_source.upper()
                    for c in model.get("columns", [])
                    if c is not bk_col
                )
                if not already_has_raw:
                    raw_row = [
                        source_schema,                              # Source Schema
                        driver_alias,                                # Source Table (alias)
                        raw_source,                                 # Source Column (stripped of cast)
                        bk_col.get("datatype", ""),                 # Datatype
                        "",                                         # Automated Logic
                        "",                                         # Manual Logic (direct passthrough)
                        "Raw source column retained alongside BK alias",  # Mapping Notes
                        order,                                      # Order#
                        raw_source,                                 # Staging Layer Column Name (stripped of cast)
                        bk_col.get("datatype", ""),                 # Staging Layer Datatype (original type)
                        "",                                         # PK
                        "",                                         # Unique
                        "",                                         # Not Null
                        "",                                         # Hashdiff
                        "",                                         # Ghost Record
                        "",                                         # Remove Column
                        "",                                         # Set Default
                        "",                                         # Test Expression
                        "",                                         # Accepted Values
                        "",                                         # Relationship
                    ]
                    ws.append(raw_row)
                    order += 1

            # Output raw BK passthrough extracted from YAML data_cols
            # (when orchestrator pre-added it — e.g., PYMNT_TERMS_CD alongside BK alias)
            #
            # Defect 3 — PREFERRED rendering: one XLSX row per source column, so a
            # composite passthrough is unambiguously machine-readable (each row has
            # a distinct staging name = the raw column). This is only possible for
            # the passthrough because its target columns ARE the raw columns.
            # (The BK-alias row above must stay a single row with newline-separated
            # source columns: it is ONE target column derived from N sources, and
            # emitting N rows would repeat the BK-alias staging name, which the
            # sheet contract forbids — _check_duplicate_staging_columns.)
            if raw_bk_col:
                hashdiff_val = str(raw_bk_col.get("hashdiff", "")).strip().lower()
                hashdiff_xlsx = "yes" if hashdiff_val in ("yes", "y") else ""
                staging_dtype = raw_bk_col.get("staging_datatype", raw_bk_col.get("datatype", ""))
                for _sc in get_source_columns(raw_bk_col):
                    raw_row = [
                        source_schema,                              # Source Schema
                        driver_alias,                               # Source Table (alias)
                        _sc,                                        # Source Column (one per row)
                        raw_bk_col.get("datatype", ""),             # Datatype
                        "",                                         # Automated Logic
                        raw_bk_col.get("manual_logic", ""),         # Manual Logic
                        "Raw source column retained alongside BK alias",  # Mapping Notes
                        order,                                      # Order#
                        _sc,                                        # Staging Layer Column Name (raw col)
                        staging_dtype,                              # Staging Datatype
                        "",                                         # PK
                        raw_bk_col.get("unique", ""),               # Unique
                        raw_bk_col.get("not_null", ""),             # Not Null
                        hashdiff_xlsx,                               # Hashdiff
                        "",                                         # Ghost Record
                        "",                                         # Remove Column
                        "",                                         # Set Default
                        "",                                         # Test Expression
                        "",                                         # Accepted Values
                        "",                                         # Relationship
                    ]
                    ws.append(raw_row)
                    order += 1

        # 2. Add HK (derived) - BEFORE data columns (requires BK to be identified)
        # Use YAML-defined HK if present, otherwise auto-generate
        if hk_cols:
            for _hk_col in hk_cols:
                hk_row = [
                    "",                                             # Source Schema (empty for derived)
                    "",                                             # Source Table
                    "(DERIVED)",                                    # Source Column
                    "BINARY",                                       # Datatype
                    "",                                             # Automated Logic
                    _hk_col.get("manual_logic", ""),                # Manual Logic from YAML
                    "Hash key aggregates BK + BKCC",                # Mapping Notes
                    order,                                          # Order#
                    _hk_col.get("staging_column_name", ""),         # Staging Layer Column Name
                    "BINARY",                                       # Staging Layer Datatype
                    "",                                             # PK
                    "",                                             # Unique
                    "",                                             # Not Null
                    "",                                             # Hashdiff
                    "",                                             # Ghost Record
                    "",                                             # Remove Column
                    "",                                             # Set Default
                    "",                                             # Test Expression
                    "",                                             # Accepted Values
                    "",                                             # Relationship
                ]
                ws.append(hk_row)
                order += 1
        elif bk_col:
            hk_name = bk_col.get("staging_column_name", "").replace("_BK", "_HK")
            if hk_name.upper() not in existing_staging_cols:
                hk_row = [
                    "",                                             # Source Schema (empty for derived)
                    "",                                             # Source Table
                    "(DERIVED)",                                    # Source Column
                    "BINARY",                                       # Datatype
                    "",                                             # Automated Logic
                    f"HASH: {source_columns_to_str(bk_col)}, BKCC",  # Manual Logic
                    "Hash key aggregates BK + BKCC",                # Mapping Notes
                    order,                                          # Order#
                    hk_name,                                        # Staging Layer Column Name
                    "BINARY",                                       # Staging Layer Datatype
                    "",                                             # PK
                    "",                                             # Unique
                    "",                                             # Not Null
                    "",                                             # Hashdiff
                    "",                                             # Ghost Record
                    "",                                             # Remove Column
                    "",                                             # Set Default
                    "",                                             # Test Expression
                    "",                                             # Accepted Values
                    "",                                             # Relationship
                ]
                ws.append(hk_row)
                order += 1

        # 3. Add data columns (in source order)
        for col in data_cols:
            hashdiff_val = str(col.get("hashdiff", "")).strip().lower()
            hashdiff_xlsx = "yes" if hashdiff_val in ("yes", "y") else ""
            src_col_list = get_source_columns(col)
            col_source_table = col.get("source_table", "")
            # Derived columns with empty source_table → keep empty (FINAL layer in build.py)
            if src_col_list == ["(DERIVED)"] and not col_source_table:
                col_alias = ""
            else:
                col_alias = col_source_table or driver_alias
            row = [
                source_schema,                              # Source Schema
                col_alias,                                   # Source Table (alias — respects secondary tables)
                _source_column_cell(col),                   # Source Column
                col.get("datatype", ""),                    # Datatype
                "",                                         # Automated Logic
                col.get("manual_logic", ""),                # Manual Logic (propagate from YAML)
                "",                                         # Mapping Notes
                order,                                      # Order#
                col.get("staging_column_name", ""),         # Staging Layer Column Name
                col.get("staging_datatype", ""),            # Staging Layer Datatype
                "",                                         # PK
                col.get("unique", ""),                      # Unique (propagate from YAML for grain cols)
                col.get("not_null", ""),                    # Not Null (propagate from YAML)
                hashdiff_xlsx,                              # Hashdiff
                "",                                         # Ghost Record
                "",                                         # Remove Column
                "",                                         # Set Default
                "",                                         # Test Expression
                "",                                         # Accepted Values
                "",                                         # Relationship
            ]
            _append_row(ws, row)
            order += 1

        # 4. Add technical columns (Fivetran, PSA)
        for col in technical_cols:
            # Technical columns: PSA_DELETE_IND and _FIVETRAN_DELETED are data for HASHDIFF
            hashdiff_val = str(col.get("hashdiff", "")).strip().lower()
            hashdiff_xlsx = "yes" if hashdiff_val in ("yes", "y") else ""
            src_col_list = get_source_columns(col)
            col_source_table = col.get("source_table", "")
            # Derived columns with empty source_table → keep empty (FINAL layer in build.py)
            if src_col_list == ["(DERIVED)"] and not col_source_table:
                col_alias = ""
            else:
                col_alias = col_source_table or driver_alias
            row = [
                source_schema,                              # Source Schema
                col_alias,                                   # Source Table (alias — respects secondary tables)
                _source_column_cell(col),                   # Source Column
                col.get("datatype", ""),                    # Datatype
                "",                                         # Automated Logic
                "",                                         # Manual Logic
                "Technical metadata column",                # Mapping Notes
                order,                                      # Order#
                col.get("staging_column_name", ""),         # Staging Layer Column Name
                col.get("staging_datatype", ""),            # Staging Layer Datatype
                "",                                         # PK
                "",                                         # Unique
                "",                                         # Not Null
                hashdiff_xlsx,                              # Hashdiff
                "",                                         # Ghost Record
                "",                                         # Remove Column
                "",                                         # Set Default
                "",                                         # Test Expression
                "",                                         # Accepted Values
                "",                                         # Relationship
            ]
            _append_row(ws, row)
            order += 1

        # 5. Add BKCC and REC_SRC (derived, from ref_business_key_collision)
        if bkcc_rec_src:
            # BKCC row — use YAML-defined if present, otherwise auto-generate
            if "BKCC" in yaml_reserved:
                yaml_bkcc = yaml_reserved["BKCC"]
                bkcc_row = [
                    "",                                             # Source Schema
                    "ref_bkcc",                                     # Source Table (must match Tables tab alias)
                    "BKCC",                                         # Source Column
                    "TEXT",                                          # Datatype
                    "",                                              # Automated Logic
                    "",                                              # Manual Logic
                    f"BKCC from ref_business_key_collision",         # Mapping Notes
                    order,                                           # Order#
                    "BKCC",                                          # Staging Layer Column Name
                    "TEXT",                                           # Staging Layer Datatype
                    "",                                              # PK
                    "",                                              # Unique
                    "",                                              # Not Null
                    "",                                              # Hashdiff
                    "bkcc",                                          # Ghost Record
                    "", "", "", "", "",                              # Remove, Set Default, Test, Values, Relationship
                ]
                ws.append(bkcc_row)
                order += 1
            elif "BKCC" not in existing_staging_cols:
                bkcc_row = [
                    "raw_vault",                                # Source Schema
                    "ref_bkcc",                                 # Source Table (alias from Tables tab)
                    "BKCC",                                     # Source Column
                    "TEXT",                                     # Datatype
                    "",                                         # Automated Logic
                    "",                                         # Manual Logic (blank - direct column)
                    f"BKCC value = {bkcc_value}",               # Mapping Notes
                    order,                                      # Order#
                    "BKCC",                                     # Staging Layer Column Name
                    "TEXT",                                     # Staging Layer Datatype
                    "",                                         # PK
                    "",                                         # Unique
                    "",                                         # Not Null
                    "",                                         # Hashdiff
                    "bkcc",                                     # Ghost Record
                    "",                                         # Remove Column
                    "",                                         # Set Default
                    "",                                         # Test Expression
                    "",                                         # Accepted Values
                    "",                                         # Relationship
                ]
                ws.append(bkcc_row)
                order += 1

            # REC_SRC row — use YAML-defined if present, otherwise auto-generate
            if "REC_SRC" in yaml_reserved:
                yaml_rec = yaml_reserved["REC_SRC"]
                rec_src_row = [
                    "",                                             # Source Schema
                    "ref_bkcc",                                     # Source Table (must match Tables tab alias)
                    "REC_SRC",                                      # Source Column
                    "TEXT",                                          # Datatype
                    "",                                              # Automated Logic
                    "",                                              # Manual Logic
                    f"REC_SRC = {bkcc_rec_src}",                    # Mapping Notes
                    order,                                           # Order#
                    "REC_SRC",                                       # Staging Layer Column Name
                    "TEXT",                                           # Staging Layer Datatype
                    "",                                              # PK
                    "",                                              # Unique
                    "",                                              # Not Null
                    "",                                              # Hashdiff
                    "rec_src",                                       # Ghost Record
                    "", "", "", "", "",                              # Remove, Set Default, Test, Values, Relationship
                ]
                ws.append(rec_src_row)
                order += 1
            elif "REC_SRC" not in existing_staging_cols:
                rec_src_row = [
                    "raw_vault",                                # Source Schema
                    "ref_bkcc",                                 # Source Table (alias from Tables tab)
                    "REC_SRC",                                  # Source Column
                    "TEXT",                                     # Datatype
                    "",                                         # Automated Logic
                    "",                                         # Manual Logic (blank - direct column)
                    f"REC_SRC = {bkcc_rec_src}",                # Mapping Notes
                    order,                                      # Order#
                    "REC_SRC",                                  # Staging Layer Column Name
                    "TEXT",                                     # Staging Layer Datatype
                    "",                                         # PK
                    "",                                         # Unique
                    "",                                         # Not Null
                    "",                                         # Hashdiff
                    "rec_src",                                  # Ghost Record
                    "",                                         # Remove Column
                    "",                                         # Set Default
                    "",                                         # Test Expression
                    "",                                         # Accepted Values
                    "",                                         # Relationship
                ]
                ws.append(rec_src_row)
                order += 1

        # 6. Add LOAD_DTS (derived)
        # Use YAML-defined if present, otherwise auto-generate
        if "LOAD_DTS" in yaml_reserved:
            yaml_load = yaml_reserved["LOAD_DTS"]
            load_dts_row = [
                "",                                             # Source Schema
                driver_alias,                                   # Source Table
                "(DERIVED)",                                    # Source Column
                "TIMESTAMP_NTZ",                                # Datatype
                "",                                             # Automated Logic
                yaml_load.get("manual_logic", "CONVERT_TIMEZONE('UTC', _FIVETRAN_SYNCED)"),  # Manual Logic from YAML
                "Load timestamp",                               # Mapping Notes
                order,                                          # Order#
                "LOAD_DTS",                                     # Staging Layer Column Name
                "TIMESTAMP_NTZ",                                # Staging Layer Datatype
                "",                                             # PK
                "",                                             # Unique
                "",                                             # Not Null
                "",                                             # Hashdiff
                "load_dts",                                     # Ghost Record
                "", "", "", "", "",                             # Remove, Set Default, Test, Values, Relationship
            ]
            ws.append(load_dts_row)
            order += 1
        elif "LOAD_DTS" not in existing_staging_cols:
            load_dts_row = [
                "",                                             # Source Schema (empty for derived)
                driver_alias,                                   # Source Table (driver alias = SRC)
                "(DERIVED)",                                    # Source Column
                "TIMESTAMP_NTZ",                                # Datatype
                "",                                             # Automated Logic
                "CONVERT_TIMEZONE('UTC', _FIVETRAN_SYNCED)",      # Manual Logic (Fivetran sources)
                "Load timestamp converted to UTC from _FIVETRAN_SYNCED (Fivetran sources)", # Mapping Notes
                order,                                          # Order#
                "LOAD_DTS",                                     # Staging Layer Column Name
                "TIMESTAMP_NTZ",                                # Staging Layer Datatype
                "",                                             # PK
                "yes",                                          # Unique (part of grain: BK + SEQUENCE_NUM + LOAD_DTS)
                "",                                             # Not Null
                "",                                             # Hashdiff
                "load_dts",                                     # Ghost Record
                "",                                             # Remove Column
                "",                                             # Set Default
                "",                                             # Test Expression
                "",                                             # Accepted Values
                "",                                             # Relationship
            ]
            ws.append(load_dts_row)
            order += 1

        # 7. Add HASHDIFF (derived, ALWAYS LAST)
        # Use YAML-defined if present, otherwise auto-generate from hashdiff=yes columns
        if "HASHDIFF" in yaml_reserved:
            yaml_hd = yaml_reserved["HASHDIFF"]
            hashdiff_row = [
                "",                                         # Source Schema
                "",                                         # Source Table
                "(DERIVED)",                                # Source Column
                "BINARY",                                   # Datatype
                "",                                         # Automated Logic
                yaml_hd.get("manual_logic", ""),            # Manual Logic from YAML
                "HASHDIFF aggregates all data columns",     # Mapping Notes
                order,                                      # Order#
                "HASHDIFF",                                 # Staging Layer Column Name
                "BINARY",                                   # Staging Layer Datatype
                "",                                         # PK
                "",                                         # Unique
                "",                                         # Not Null
                "",                                         # Hashdiff
                "hashdiff",                                 # Ghost Record
                "", "", "", "", "",                         # Remove, Set Default, Test, Values, Relationship
            ]
            ws.append(hashdiff_row)
        elif "HASHDIFF" not in existing_staging_cols:
            hashdiff_cols = [col.get("staging_column_name", "") for col in model.get("columns", []) if col.get("hashdiff") == "yes"]
            if hashdiff_cols:
                hashdiff_row = [
                    "",                                         # Source Schema (empty for derived)
                    "",                                         # Source Table
                    "(DERIVED)",                                # Source Column
                    "BINARY",                                   # Datatype
                    "",                                         # Automated Logic
                    f"HASH: {', '.join(hashdiff_cols)}",        # Manual Logic
                    "HASHDIFF aggregates all data columns (not BK, metadata, or technical)",  # Mapping Notes
                    order,                                      # Order#
                    "HASHDIFF",                                 # Staging Layer Column Name
                    "BINARY",                                   # Staging Layer Datatype
                    "",                                         # PK
                    "",                                         # Unique
                    "",                                         # Not Null
                    "",                                         # Hashdiff
                    "hashdiff",                                 # Ghost Record
                    "",                                         # Remove Column
                    "",                                         # Set Default
                    "",                                         # Test Expression
                    "",                                         # Accepted Values
                    "",                                         # Relationship
                ]
                ws.append(hashdiff_row)

    def _make_column_row(self, col, order, layer="STG"):
        base = [
            col.get("source_schema", ""),
            col.get("source_table", ""),
            _source_column_cell(col),
            col.get("datatype", ""),
            col.get("automated_logic", ""),
            col.get("manual_logic", ""),
            col.get("mapping_notes", ""),
            col.get("order", order),
            col.get("staging_column_name", ""),
            col.get("staging_datatype", ""),
            col.get("pk", ""),
            col.get("unique", ""),
        ]

        layer_upper = layer.upper()
        if layer_upper in ("HUB", "LNK"):
            # HUB/LNK: PK, Unique, Not Null, Ghost Record, Hashdiff
            base += [
                col.get("not_null", ""),
                col.get("ghost_record", ""),
                col.get("hashdiff", ""),
            ]
        elif layer_upper in ("SAT", "LSAT", "MSAT", "LMSAT", "ESAT"):
            # SAT: PK, Unique, Ghost Record, Not Null (no Hashdiff header)
            base += [
                col.get("ghost_record", ""),
                col.get("not_null", ""),
            ]
        else:
            # STG: PK, Unique, Not Null, Hashdiff, Ghost Record
            base += [
                col.get("not_null", ""),
                col.get("hashdiff", ""),
                col.get("ghost_record", ""),
            ]

        base += [
            col.get("remove_column", ""),
            col.get("set_default", ""),
            col.get("test_expression", ""),
            col.get("accepted_values", ""),
            col.get("relationship", ""),
        ]
        return base

    def _make_meta_row(self, source_table, col_name, dtype, ghost_val, order, layer="STG"):
        """Build a meta column row with correct field ordering for the layer."""
        base = [
            "",             # Source Schema
            source_table,   # Source Table (alias)
            col_name,       # Source Column
            dtype,          # Datatype
            "",             # Automated Logic
            "",             # Manual Logic
            "",             # Mapping Notes
            order,          # Order#
            col_name,       # Staging Layer Column Name
            dtype,          # Staging Layer Datatype
            "",             # PK
            "",             # Unique
        ]

        layer_upper = layer.upper()
        if layer_upper in ("HUB", "LNK"):
            base += ["", ghost_val, ""]   # Not Null, Ghost Record, Hashdiff
        elif layer_upper in ("SAT", "LSAT", "MSAT", "LMSAT", "ESAT"):
            base += [ghost_val, ""]       # Ghost Record, Not Null
        else:
            base += ["", "", ghost_val]   # Not Null, Hashdiff, Ghost Record

        base += ["", "", "", "", ""]  # Remove Column, Set Default, Test Expr, Values, Relationship
        return base


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Generate tech-spec XLSX for process-mapping")
    parser.add_argument("--config", required=True, help="Path to config YAML defining the models")
    parser.add_argument("--outdir", required=True, help="Output directory (should be <rootdir>/mappings/)")
    parser.add_argument("--filename", default=None, help="Output XLSX filename (default: from config)")
    args = parser.parse_args()

    with open(args.config, "r") as f:
        config = yaml.safe_load(f)

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    filename = args.filename or config.get("filename", "tech_spec.xlsx")
    if not filename.endswith(".xlsx"):
        filename += ".xlsx"

    builder = TechSpecBuilder(config)
    wb = builder.build()

    outpath = outdir / filename
    wb.save(str(outpath))
    print(f"Created: {outpath}")
    print(f"Models defined: {len(config.get('models', []))}")
    print(f"\nNext step: python src/main.py --rootdir {outdir.parent} --debug")


if __name__ == "__main__":
    main()
